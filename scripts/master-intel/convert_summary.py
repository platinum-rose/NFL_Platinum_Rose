# C:\Users\andre\.gemini\antigravity\brain\b627786e-a938-479a-b576-b3b6c7c2b993\scratch\convert_summary.py
import re
import sys
from pathlib import Path
try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
except ImportError:
    docx = None
    Pt = lambda x: x
    Inches = lambda x: x
    RGBColor = lambda *args: None
    WD_TABLE_ALIGNMENT = type('WD_TABLE_ALIGNMENT', (), {'CENTER': 1})
    WD_ALIGN_VERTICAL = type('WD_ALIGN_VERTICAL', (), {'CENTER': 1})
    parse_xml = lambda x: None
    nsdecls = lambda x: ''

CITE_CSS = r'''
.glance-group-label{font-size:12px;font-weight:700;letter-spacing:.06em;text-transform:uppercase;color:var(--muted);margin:10px 0 4px}.glance-group-label span{font-weight:400}

/* ===== game pop-up: header, bullets, camps, secondary callout ===== */
.gh { display: grid; grid-template-columns: auto 1fr auto; align-items: center; gap: 16px; margin: 6px 0 12px; }
.gh-team { display: flex; flex-direction: column; align-items: center; gap: 2px; min-width: 96px; text-align: center; }
.gh-team b { font-size: 20px; color: var(--text); letter-spacing: .02em; }
.gh-team small { color: var(--muted); font-size: 11.5px; }
.gh-mid { min-width: 0; }
.gh table.mkt-lines { width: min(520px, 100%); margin: 0 auto; }
.gs-list { margin: 4px 0 6px; padding-left: 20px; }
.gs-list li { margin: 4px 0; line-height: 1.5; }
.sec-callout { margin: 6px 0 14px; padding: 12px 14px; border: 1px solid #f59e0b; border-left-width: 5px; border-radius: 10px; background: rgba(245, 158, 11, 0.09); }
.sec-title { font-weight: 800; color: var(--text); font-size: 15px; margin-bottom: 8px; }
.sec-title span { display: block; font-weight: 500; font-size: 12px; color: var(--muted); margin-top: 1px; }
.sec-row { display: flex; gap: 10px; align-items: flex-start; padding: 7px 0; border-top: 1px solid rgba(245, 158, 11, .25); }
.sec-badge { flex: none; padding: 3px 9px; border-radius: 999px; font-size: 11.5px; font-weight: 800; background: #f59e0b; color: #0f172a; white-space: nowrap; }
.sec-medium .sec-badge { background: #facc15; }
.sec-targets { margin-top: 4px; font-size: 12.5px; color: var(--muted); }
.sec-chip { display: inline-block; margin: 2px 4px 0 4px; padding: 1px 9px; border-radius: 999px; background: var(--card); border: 1px solid var(--border); color: var(--text); font-weight: 700; }
.sec-note { margin-top: 6px; font-size: 11.5px; color: var(--muted); }
.gp-list { list-style: none; margin: 4px 0; padding: 0; }
.gp-list li { margin: 5px 0; padding: 6px 10px; border-radius: 8px; background: var(--raised); line-height: 1.45; font-size: 13px; }
.gp-list li.gp-ours { border-left: 3px solid var(--accent); }
.gp-list li.gp-none { color: var(--muted); background: transparent; }
.gp-src { color: var(--muted); }
.gp-why { color: var(--body); }
.sent-head { display: grid; grid-template-columns: auto 1fr auto; align-items: center; gap: 10px; margin: 4px 0 8px; font-size: 13px; color: var(--text); }
.sent-bar { position: relative; display: block; height: 10px; border-radius: 6px; background: #f59e0b; overflow: hidden; }
.sent-bar i { position: absolute; left: 0; top: 0; bottom: 0; background: var(--accent); }
.sent-cols { display: grid; grid-template-columns: 1fr 1fr; gap: 12px; }
.sent-col-t { font-weight: 800; font-size: 12px; text-transform: uppercase; letter-spacing: .04em; color: var(--muted); margin-bottom: 2px; }
.sent-other { margin-top: 8px; }
@media (max-width: 640px) {
  .gh { grid-template-columns: auto auto; grid-template-areas: "a h" "m m"; justify-content: space-between; }
  .gh > .gh-team:first-child { grid-area: a; } .gh > .gh-team:last-child { grid-area: h; } .gh-mid { grid-area: m; }
  .gh .tl-xl { width: 56px !important; height: 56px !important; }
  .sent-cols { grid-template-columns: 1fr; }
}

.mkt-top { display: flex; justify-content: center; margin: 4px 0 10px; }
table.mkt-lines { width: min(560px, 100%); margin: 0 auto; border-collapse: collapse; text-align: center; font-size: 13.5px; }
table.mkt-lines th, table.mkt-lines td { padding: 7px 10px; border-bottom: 1px solid var(--border); text-align: center; }
table.mkt-lines thead th { color: var(--text); font-weight: 800; }
table.mkt-lines tbody th { text-align: left; color: var(--muted); font-weight: 700; white-space: nowrap; }

.spl-cell { display: block; padding: 8px 10px; border-radius: 8px; background: var(--card); }
.spl-cell b { display: block; margin-bottom: 4px; color: var(--text); font-size: 11px; text-transform: uppercase; letter-spacing: .04em; }
.spl-row { display: grid; grid-template-columns: 52px 1fr; grid-template-rows: auto auto; column-gap: 8px; row-gap: 2px; padding: 4px 6px; margin: 2px 0; border-radius: 6px; }
.spl-side { grid-row: 1 / span 2; align-self: center; font-weight: 800; color: var(--text); }
.spl-big { background: rgba(250, 204, 21, 0.12); box-shadow: inset 3px 0 0 #facc15; }
.spl-bar { position: relative; display: block; height: 14px; border-radius: 4px; background: var(--border); overflow: hidden; }
.spl-bar i { position: absolute; left: 0; top: 0; bottom: 0; display: block; border-radius: 4px; }
.spl-bar em { position: absolute; right: 6px; top: 0; font-style: normal; font-size: 10.5px; line-height: 14px; font-weight: 800; color: var(--text); text-shadow: 0 0 3px rgba(0,0,0,.6); }
.spl-bets i, .spl-kb { background: #60a5fa; }
.spl-money i, .spl-km { background: var(--accent); }
.spl-key { display: flex; flex-wrap: wrap; align-items: center; gap: 6px; margin-bottom: 8px; font-size: 11.5px; color: var(--muted); }
.spl-k { display: inline-block; width: 12px; height: 8px; border-radius: 2px; margin-left: 4px; }
.spl-note { margin-left: 8px; }

/* ===== betting splits strip and movement badges ===== */
.spl-strip { border: 1px solid var(--border); border-radius: 10px; background: var(--raised); padding: 10px 12px; margin: 4px 0 14px; font-size: 12.5px; }
.spl-head { display: flex; flex-wrap: wrap; gap: 8px; align-items: center; margin-bottom: 6px; }
.spl-when { color: var(--muted); font-size: 11.5px; }
.spl-cells { display: grid; grid-template-columns: repeat(auto-fit, minmax(210px, 1fr)); gap: 8px; }
.spl-cell { display: block; padding: 6px 8px; border-radius: 8px; background: var(--card); line-height: 1.45; }
.spl-cell b { display: block; color: var(--text); font-size: 11px; text-transform: uppercase; letter-spacing: .04em; }
.spl-mv { margin-top: 6px; color: var(--muted); }
.mv-badge { display: inline-block; padding: 1px 8px; border-radius: 999px; font-size: 11.5px; font-weight: 800; border: 1px solid var(--border); background: var(--card); color: var(--text); cursor: help; white-space: nowrap; }
.mv-steam { border-color: #f59e0b; color: #f59e0b; }
.mv-rlm { border-color: #a78bfa; color: #a78bfa; }
.mv-bigmove { border-color: var(--accent); color: var(--accent); }

/* ===== citation popups + glossary hovers ===== */
.cite, .gl { position: relative; cursor: help; }
.gl { border-bottom: 1px dotted var(--accent); text-decoration: none; }
.cite { border-bottom: 1px dotted var(--accent); font-weight: 700; }
.cite-pop { display: none; position: absolute; left: 0; top: 100%; z-index: 60; width: min(340px, 86vw); padding: 10px 12px; margin-top: 2px; border: 1px solid var(--border); border-radius: 10px; background: var(--card); color: var(--body); font-size: 12.5px; font-weight: 400; line-height: 1.4; box-shadow: 0 8px 24px rgba(0,0,0,.4); text-align: left; cursor: default; }
.cite:hover .cite-pop, .cite:focus .cite-pop, .cite:focus-within .cite-pop { display: block; }
.cite-pop b { color: var(--text); }
.cite-row { display: block; margin-top: 8px; }
.cite-row i { display: block; margin: 2px 0; color: var(--muted); }
.cite-row a, .cite-all { display: inline-block; margin-top: 2px; color: var(--accent); font-weight: 700; }
.cite-row em { display: block; color: var(--muted); font-size: 11.5px; }
.cite-all { margin-top: 8px; }
.gl:hover::after, .gl:focus::after { content: attr(data-tip); position: absolute; left: 0; top: 100%; z-index: 60; width: min(300px, 80vw); padding: 8px 10px; margin-top: 4px; border: 1px solid var(--border); border-radius: 8px; background: var(--card); color: var(--body); font-size: 12.5px; font-weight: 400; line-height: 1.4; box-shadow: 0 8px 24px rgba(0,0,0,.4); white-space: normal; }
@media print { .cite-pop, .gl::after { display: none !important; } }
'''

MT_CSS = r'''/* ===== Master Intel tabbed layout (2026-10-17). Active only when body.mt-on is set by initMasterTabs(). ===== */
html { scroll-behavior: auto !important; }
:target { animation: none !important; }
[id] { scroll-margin-top: 76px; }
.target-pulse { animation: none !important; outline: 2px solid var(--accent); outline-offset: 2px; border-radius: 6px; }
body.mt-on { padding: 0 0 28px; }
body.mt-on .container { margin-top: 14px; border: 0; padding: 6px 20px 24px; }
body.mt-on .container > h1,
body.mt-on .container > h2,
body.mt-on .global-rollup-bar,
body.mt-on .toc-title,
body.mt-on .toc-grid,
body.mt-on .part-banner,
body.mt-on .side-toc,
body.mt-on .toc-toggle { display: none !important; }
body.mt-on .container > h3 { font-size: 13px; font-weight: 500; color: var(--muted); margin: 4px 0 10px; }
.mt-bar { position: sticky; top: 0; z-index: 1000; display: flex; align-items: center; gap: 18px; min-height: 56px; padding: 0 20px; background: var(--card); border-bottom: 1px solid var(--border); }
.mt-brand { display: flex; align-items: center; gap: 8px; font-weight: 800; color: var(--text); white-space: nowrap; font-size: 15px; }
.mt-tabs { display: flex; gap: 4px; flex: 1; overflow-x: auto; scrollbar-width: none; }
.mt-tabs::-webkit-scrollbar { display: none; }
.mt-tab { min-height: 44px; padding: 0 16px; border: 0; border-radius: 10px; background: transparent; color: var(--body); font: inherit; font-weight: 700; font-size: 14px; cursor: pointer; display: inline-flex; align-items: center; gap: 7px; white-space: nowrap; }
.mt-tab:hover { background: var(--raised); color: var(--text); }
.mt-tab[aria-selected="true"] { background: var(--accent); color: var(--on-accent); }
.mt-ico { font-size: 15px; line-height: 1; }
.mt-panel { display: none; }
.mt-panel.active { display: block; }
.mt-panels { min-height: 70vh; }
.mt-panel > h2 { margin-top: 26px; }
.mt-panel > h2:first-child, .mt-panel > a:first-child + h2 { margin-top: 8px; }
.mt-section { margin-bottom: 18px; }
.mt-bottom { display: none; position: fixed; left: 0; right: 0; bottom: 0; z-index: 1000; background: var(--card); border-top: 1px solid var(--border); padding-bottom: env(safe-area-inset-bottom); }
.mt-bt { flex: 1; min-height: 58px; display: flex; flex-direction: column; align-items: center; justify-content: center; gap: 3px; border: 0; background: transparent; color: var(--muted); font: inherit; font-size: 11px; font-weight: 700; cursor: pointer; }
.mt-bt .mt-ico { font-size: 21px; }
.mt-bt.active { color: var(--accent); box-shadow: inset 0 3px 0 var(--accent); }
.mt-sheet { position: fixed; inset: 0; z-index: 1100; background: rgba(2, 6, 23, 0.6); display: flex; align-items: flex-end; }
.mt-sheet[hidden] { display: none; }
.mt-sheet-card { width: 100%; background: var(--card); border-top: 1px solid var(--border); border-radius: 18px 18px 0 0; padding: 14px 14px calc(18px + env(safe-area-inset-bottom)); }
.mt-sheet-title { font-weight: 800; color: var(--text); margin: 2px 4px 10px; }
.mt-sheet-btn { width: 100%; min-height: 52px; display: flex; align-items: center; gap: 12px; margin-bottom: 8px; padding: 0 14px; border: 1px solid var(--border); border-radius: 12px; background: var(--raised); color: var(--text); font: inherit; font-weight: 700; font-size: 15px; cursor: pointer; text-align: left; }
.mt-sheet-btn.active { background: var(--accent); color: var(--on-accent); border-color: var(--accent); }
.mt-sheet-close { width: 100%; min-height: 46px; margin-top: 2px; border: 0; background: transparent; color: var(--muted); font: inherit; font-weight: 700; cursor: pointer; }
/* game grid */
.gg-bar { display: flex; flex-wrap: wrap; align-items: center; justify-content: space-between; gap: 10px; margin: 6px 0 12px; }
.gg-chips { display: flex; gap: 6px; flex-wrap: wrap; }
.gg-chip { min-height: 40px; padding: 0 14px; border: 1px solid var(--border); border-radius: 999px; background: var(--raised); color: var(--body); font: inherit; font-weight: 700; font-size: 13px; cursor: pointer; }
.gg-chip.active { background: var(--accent); color: var(--on-accent); border-color: var(--accent); }
.gg-sort { font-size: 13px; color: var(--muted); display: inline-flex; align-items: center; gap: 6px; }
.gg-sort select { min-height: 40px; padding: 0 8px; border: 1px solid var(--border); border-radius: 8px; background: var(--raised); color: var(--text); font: inherit; }
.gg-cards { display: grid; grid-template-columns: repeat(auto-fill, minmax(320px, 1fr)); gap: 12px; }
.gg-card { display: block; padding: 12px 14px; border: 1px solid var(--border); border-radius: 12px; background: var(--card); color: var(--body); text-decoration: none; }
.gg-card:hover, .gg-card:focus-visible { border-color: var(--accent); }
.gg-card[hidden] { display: none; }
.gg-final { opacity: .7; }
.gg-top { display: flex; justify-content: space-between; gap: 8px; font-size: 12px; color: var(--muted); font-weight: 700; }
.gg-badge-final { color: var(--text); }
.gg-teams { font-weight: 800; color: var(--text); font-size: 16px; margin: 4px 0 2px; }
.gg-line { font-size: 12.5px; color: var(--muted); margin-bottom: 8px; }
.gg-pick { display: flex; align-items: baseline; gap: 8px; padding: 6px 0; border-top: 1px solid var(--border); }
.gg-lbl { font-size: 11px; text-transform: uppercase; color: var(--muted); font-weight: 700; min-width: 40px; }
.gg-pick b { color: var(--text); flex: 1; }
.gg-px { color: var(--muted); font-size: 13px; }
.gg-conf { font-weight: 800; color: var(--accent); font-size: 14px; }
.gg-conf small { font-weight: 600; color: var(--muted); margin-left: 3px; font-size: 10px; }
.gg-risk { font-size: 12px; color: var(--muted); margin-top: 6px; display: -webkit-box; -webkit-line-clamp: 2; -webkit-box-orient: vertical; overflow: hidden; }
.gg-note { font-size: 12px; color: var(--muted); margin: 10px 0 0; }
.gg-nav { display: none; align-items: center; gap: 8px; margin: 6px 0 12px; flex-wrap: wrap; }
.gg-nav button { min-height: 44px; padding: 0 14px; border: 1px solid var(--border); border-radius: 10px; background: var(--raised); color: var(--text); font: inherit; font-weight: 700; cursor: pointer; }
.gg-nav .gg-spacer { flex: 1; }
.mt-panel.gg-detail .gg-nav { display: flex; position: sticky; top: 56px; z-index: 50; background: var(--bg, var(--card)); padding: 6px 0; }
.gg-off { display: none !important; }
.mt-panel.gg-detail #gg-grid { display: none; }
@media (max-width: 640px) {
  .gg-cards { grid-template-columns: 1fr; }
  body.mt-on { padding-bottom: calc(76px + env(safe-area-inset-bottom)); }
  body.mt-on .mt-tabs { display: none; }
  body.mt-on .mt-bar { min-height: 46px; padding: 0 12px; }
  body.mt-on .mt-bottom { display: flex; }
  body.mt-on .container { padding: 4px 10px 20px; }
  body.mt-on .card-builder-trigger { bottom: calc(76px + env(safe-area-inset-bottom)); right: 12px; }
  body.mt-on .back-to-top { bottom: calc(132px + env(safe-area-inset-bottom)); }
}
@media (prefers-reduced-motion: reduce) { .mt-tab, .mt-bt { transition: none; } }
@media print {
  .mt-bar, .mt-bottom, .mt-sheet { display: none !important; }
  .mt-panel { display: block !important; }
}

.mt-jump { margin: 4px 0 12px; font-size: 14px; font-weight: 700; }
.mt-jump a { color: var(--accent); }
'''

MT_JS = r'''function initMasterTabs() {
  if (window.__mtInit) return;
  const c = document.querySelector('.container');
  if (!c || c.parentNode !== document.body) return;
  if (document.getElementById('sc-our-five')) return;
  if (!document.getElementById('section-rec-box') || !document.getElementById('section-1-box') || !document.getElementById('section-6-box') || !document.getElementById('section-9-box')) return;
  window.__mtInit = true;
  const TABS = [
    { id: 'picks', label: 'Picks', icon: '\u{1F3AF}' },
    { id: 'games', label: 'Games', icon: '\u{1F3C8}' },
    { id: 'props', label: 'Props', icon: '\u{1F9FE}' },
    { id: 'experts', label: 'Experts', icon: '\u{1F399}️' },
    { id: 'specials', label: 'Specials', icon: '\u{1F39F}️' },
    { id: 'market', label: 'Market', icon: '\u{1F4C8}' },
    { id: 'about', label: 'About', icon: '\u{1F4DA}' }
  ];
  const BOTTOM = ['picks', 'games', 'props', 'experts'];
  const txt = el => (el.textContent || '').trim();
  function marker(el) {
    const id = el.id || '';
    if (id === 'recommendations') return ['picks', 1];
    if (id === 'executive-summary') return ['picks', 2];
    if (el.tagName === 'H2' && /^1\./.test(txt(el))) return ['picks', 3];
    if (id === 'prop-card') return ['props', 1];
    if (id === 'teaser-matrix') return ['specials', 1];
    if (id === 'game-dossier') return ['games', 1];
    if (id === 'consensus-section') return ['experts', 1];
    if (el.tagName === 'H2' && /^8\./.test(txt(el))) return ['experts', 2];
    if (id === 'systems') return ['market', 1];
    if (id === 'how-card-built') return ['about', 1];
    return null;
  }
  const kids = Array.from(c.children);
  const groups = {};
  let cur = null;
  kids.forEach(k => {
    const m = marker(k);
    if (m) { cur = { rank: m[1], nodes: [] }; (groups[m[0]] = groups[m[0]] || []).push(cur); }
    if (cur) cur.nodes.push(k);
  });
  const wrap = document.createElement('div');
  wrap.className = 'mt-panels';
  const panels = {};
  TABS.forEach(t => {
    const p = document.createElement('section');
    p.className = 'mt-panel'; p.id = 'mt-panel-' + t.id;
    p.setAttribute('role', 'tabpanel'); p.setAttribute('aria-labelledby', 'mt-tab-' + t.id);
    p.setAttribute('inert', '');
    (groups[t.id] || []).sort((a, b) => a.rank - b.rank).forEach(g => g.nodes.forEach(n => p.appendChild(n)));
    panels[t.id] = p; wrap.appendChild(p);
  });
  c.appendChild(wrap);
  const slate = document.getElementById('slate-status-box');
  if (slate) panels.picks.insertBefore(slate, panels.picks.firstChild);
  // dissolve the outer "Section N" accordions: keep content and the id (old links still land), drop the box
  wrap.querySelectorAll('details.section-main').forEach(d => {
    const rc = d.querySelector(':scope > .rollup-content');
    const w = document.createElement('div');
    w.className = 'mt-section'; w.id = d.id;
    d.classList.forEach(cl => { if (cl === 'section-main' || /^secmain-/.test(cl)) w.classList.add(cl); });
    if (rc) while (rc.firstChild) w.appendChild(rc.firstChild);
    d.parentNode.replaceChild(w, d);
  });
  // "Every expert's picks this week" belongs with the experts, not the picks
  (function () {
    const a = document.getElementById('expert-best-bets');
    if (!a || !a.parentNode) return;
    const box = document.createElement('div');
    box.id = 'mt-expert-picks'; box.className = 'mt-section';
    let n = a;
    while (n) {
      const nx = n.nextElementSibling;
      const ok = n === a || n.tagName === 'H3' || n.tagName === 'P' || (n.tagName === 'DETAILS' && /^bx-/.test(n.id || ''));
      if (!ok) break;
      box.appendChild(n); n = nx;
    }
    panels.experts.appendChild(box);
    const jump = document.createElement('p');
    jump.className = 'mt-jump';
    jump.innerHTML = '<a href="#expert-best-bets">\u2B50 Every expert\'s picks this week, by name \u2193</a>';
    panels.experts.insertBefore(jump, panels.experts.firstChild);
  })();
  // the market board (line movement, steam/RLM, splits, all-games table) belongs on the Market tab
  (function () {
    const box = document.createElement('div');
    box.id = 'mt-market-board'; box.className = 'mt-section';
    ['line-moves', 'platinum-rose-board', 'board-defs'].forEach(id => { const e = document.getElementById(id); if (e) box.appendChild(e); });
    if (box.children.length) panels.market.insertBefore(box, panels.market.firstChild);
  })();
  // outer Expand/Collapse Section buttons have nothing left to act on
  wrap.querySelectorAll(':scope > section > .rollup-controls, .mt-panel > .rollup-controls').forEach(rc => {
    rc.querySelectorAll('button').forEach(b => { if (/secmain-/.test(b.getAttribute('onclick') || '')) b.remove(); });
    if (!rc.querySelector('button')) rc.remove();
  });
  wrap.querySelectorAll('a[href="#recommendations"]').forEach(a => {
    if (/Back to/i.test(txt(a))) { a.textContent = '⬆ Back to top'; a.setAttribute('href', '#mt-top'); }
  });
  const side = document.getElementById('side-toc'); if (side) side.remove();
  const tocb = document.querySelector('.toc-toggle'); if (tocb) tocb.remove();
  document.body.classList.add('mt-on');

  const wk = document.body.getAttribute('data-week') || '';
  const bar = document.createElement('header');
  bar.className = 'mt-bar'; bar.id = 'mt-bar';
  bar.innerHTML = '<div class="mt-brand"><span aria-hidden="true">\u{1F3C8}</span><span>Week ' + wk + ' Master Intel</span></div><div class="mt-tabs" role="tablist" aria-label="Report sections"></div>';
  const tl = bar.querySelector('.mt-tabs');
  const tabBtns = {};
  TABS.forEach(t => {
    const b = document.createElement('button');
    b.type = 'button'; b.className = 'mt-tab'; b.id = 'mt-tab-' + t.id; b.dataset.tab = t.id;
    b.setAttribute('role', 'tab'); b.setAttribute('aria-controls', 'mt-panel-' + t.id); b.setAttribute('aria-selected', 'false'); b.tabIndex = -1;
    b.innerHTML = '<span class="mt-ico" aria-hidden="true">' + t.icon + '</span><span>' + t.label + '</span>';
    tl.appendChild(b); tabBtns[t.id] = b;
  });
  const top = document.createElement('span'); top.id = 'mt-top';
  document.body.insertBefore(top, document.body.firstChild);
  document.body.insertBefore(bar, c);

  const bn = document.createElement('nav');
  bn.className = 'mt-bottom'; bn.setAttribute('aria-label', 'Report sections');
  const btmBtns = {};
  BOTTOM.forEach(id => {
    const t = TABS.find(x => x.id === id);
    const b = document.createElement('button');
    b.type = 'button'; b.className = 'mt-bt'; b.dataset.tab = id;
    b.innerHTML = '<span class="mt-ico" aria-hidden="true">' + t.icon + '</span><span>' + t.label + '</span>';
    bn.appendChild(b); btmBtns[id] = b;
  });
  const more = document.createElement('button');
  more.type = 'button'; more.className = 'mt-bt mt-more'; more.setAttribute('aria-haspopup', 'dialog'); more.setAttribute('aria-expanded', 'false');
  more.innerHTML = '<span class="mt-ico" aria-hidden="true">⋯</span><span>More</span>';
  bn.appendChild(more);
  document.body.appendChild(bn);
  const sheet = document.createElement('div');
  sheet.className = 'mt-sheet'; sheet.id = 'mt-sheet'; sheet.hidden = true;
  sheet.setAttribute('role', 'dialog'); sheet.setAttribute('aria-modal', 'true'); sheet.setAttribute('aria-label', 'More sections');
  let sh = '<div class="mt-sheet-card"><div class="mt-sheet-title">More</div>';
  TABS.filter(t => BOTTOM.indexOf(t.id) < 0).forEach(t => { sh += '<button type="button" class="mt-sheet-btn" data-tab="' + t.id + '"><span class="mt-ico" aria-hidden="true">' + t.icon + '</span><span>' + t.label + '</span></button>'; });
  sh += '<button type="button" class="mt-sheet-close">Close</button></div>';
  sheet.innerHTML = sh;
  document.body.appendChild(sheet);

  let current = null;
  const scrollPos = {};
  const barH = () => bar.offsetHeight || 56;
  function closeSheet() { sheet.hidden = true; more.setAttribute('aria-expanded', 'false'); }
  function show(id, o) {
    o = o || {};
    if (!panels[id]) id = 'picks';
    if (current) scrollPos[current] = window.scrollY;
    const was = current;
    TABS.forEach(t => {
      const on = t.id === id, p = panels[t.id];
      p.classList.toggle('active', on);
      if (on) p.removeAttribute('inert'); else p.setAttribute('inert', '');
      tabBtns[t.id].setAttribute('aria-selected', on ? 'true' : 'false');
      tabBtns[t.id].tabIndex = on ? 0 : -1;
      if (btmBtns[t.id]) btmBtns[t.id].classList.toggle('active', on);
    });
    more.classList.toggle('active', BOTTOM.indexOf(id) < 0);
    sheet.querySelectorAll('.mt-sheet-btn').forEach(b => b.classList.toggle('active', b.dataset.tab === id));
    current = id;
    closeSheet();
    try { sessionStorage.setItem('mt-tab', id); } catch (e) {}
    const b = tabBtns[id];
    if (b && tl.scrollWidth > tl.clientWidth) tl.scrollLeft = b.offsetLeft - (tl.clientWidth - b.offsetWidth) / 2;
    if (o.scroll !== false) {
      const y = o.top != null ? o.top : (was === id ? 0 : (scrollPos[id] || 0));
      window.scrollTo(0, y);
    }
    if (o.push !== false) {
      const h = '#tab-' + id;
      if (location.hash !== h) { try { history.pushState(null, '', h); } catch (e) {} }
    }
  }
  function reveal(hash) {
    const raw = (hash || '').replace(/^#/, '');
    if (!raw) return false;
    if (raw === 'mt-top') { window.scrollTo(0, 0); return true; }
    const mt = raw.match(/^tab-([a-z]+)$/);
    if (mt) { show(mt[1], { push: false }); if (mt[1] === 'games' && typeof closeGame === 'function') closeGame(false); return true; }
    const el = document.getElementById(raw);
    if (!el) return false;
    const p = el.closest('.mt-panel');
    if (!p) return false;
    show(p.id.replace('mt-panel-', ''), { push: false, scroll: false });
    const gb = el.closest('[id^="game-"][id$="-box"]') || (/^game-.*-box$/.test(raw) ? el : null);
    if (gb && p.id === 'mt-panel-games') { openGame(gb, false); window.scrollTo(0, 0); return true; }
    else if (p.id === 'mt-panel-games' && !el.closest('[id^="game-"]')) { closeGame(false); }
    let n = el;
    while (n && n !== document.body) { if (n.tagName === 'DETAILS') n.open = true; n = n.parentElement; }
    requestAnimationFrame(() => {
      const y = el.getBoundingClientRect().top + window.scrollY - barH() - 8;
      window.scrollTo(0, Math.max(0, y));
    });
    el.classList.remove('target-pulse'); void el.offsetWidth; el.classList.add('target-pulse');
    setTimeout(() => el.classList.remove('target-pulse'), 1800);
    return true;
  }
  // ---- Games tab: grid of cards <-> one game at a time ----
  const gp = panels.games;
  const gboxes = Array.from(gp.querySelectorAll('[id^="game-"][id$="-box"]')).filter(x => x.closest('.mt-panel') === gp);
  const nav = document.createElement('div');
  nav.className = 'gg-nav';
  nav.innerHTML = '<button type="button" data-g="back">\u2190 All games</button><span class="gg-spacer"></span><button type="button" data-g="prev">\u2039 Prev</button><button type="button" data-g="next">Next \u203A</button>';
  const grid = document.getElementById('gg-grid');
  let openBox = null;
  function gcards() { return Array.from(document.querySelectorAll('#gg-grid .gg-card')); }
  function applyGrid() {
    const f = (document.querySelector('.gg-chip.active') || {}).dataset; const day = f ? f.f : 'all';
    const sel = document.getElementById('gg-sort'); const mode = sel ? sel.value : 'kick';
    const cs = gcards();
    cs.forEach(c => { c.hidden = !(day === 'all' || c.dataset.day === day); });
    const parent = cs[0] && cs[0].parentNode;
    if (parent) cs.slice().sort((a, b) => {
      const fa = a.classList.contains('gg-final') ? 1 : 0, fb = b.classList.contains('gg-final') ? 1 : 0;
      if (fa !== fb) return fa - fb;
      return mode === 'conf' ? (parseFloat(b.dataset.conf) - parseFloat(a.dataset.conf)) || (a.dataset.kick - b.dataset.kick) : a.dataset.kick - b.dataset.kick;
    }).forEach(c => parent.appendChild(c));
  }
  function openGame(box, push) {
    if (!box) return;
    gboxes.forEach(b => { b.hidden = b !== box; });
    box.hidden = false;
    if (box.tagName === 'DETAILS') box.open = true;
    openBox = box;
    gp.classList.add('gg-detail');
    gp.querySelectorAll('.gg-off').forEach(x => x.classList.remove('gg-off'));
    for (let n = box; n && n !== gp; n = n.parentNode) Array.from(n.parentNode.children).forEach(s => { if (s !== n && s !== nav) s.classList.add('gg-off'); });
    const gids = gboxes.map(b => b.id);
    const i = gids.indexOf(box.id);
    nav.querySelector('[data-g=prev]').disabled = i <= 0;
    nav.querySelector('[data-g=next]').disabled = i >= gids.length - 1;
    if (push !== false) { try { history.pushState(null, '', '#' + box.id); } catch (e) {} window.scrollTo(0, 0); }
  }
  function closeGame(push) {
    gp.classList.remove('gg-detail'); openBox = null;
    gp.querySelectorAll('.gg-off').forEach(x => x.classList.remove('gg-off'));
    gboxes.forEach(b => { b.hidden = true; });
    if (push !== false) { try { history.pushState(null, '', '#tab-games'); } catch (e) {} window.scrollTo(0, 0); }
  }
  if (grid && gboxes.length) {
    const h6 = Array.from(gp.querySelectorAll('h2')).find(h => /^6\./.test((h.textContent || '').trim()));
    if (h6 && h6.parentNode) h6.parentNode.insertBefore(grid, h6.nextSibling);
    gp.insertBefore(nav, gp.firstChild);
    gboxes.forEach(b => { b.hidden = true; });
    const gh = gp.querySelector(':scope > .rollup-controls'); if (gh) gh.style.display = 'none';
    grid.addEventListener('click', e => {
      const ch = e.target.closest('.gg-chip');
      if (ch) { grid.querySelectorAll('.gg-chip').forEach(x => x.classList.toggle('active', x === ch)); applyGrid(); return; }
      const c = e.target.closest('.gg-card');
      if (c) { e.preventDefault(); const bx = document.getElementById(c.getAttribute('href').slice(1)); if (bx) openGame(bx); }
    });
    grid.addEventListener('change', e => { if (e.target.id === 'gg-sort') applyGrid(); });
    nav.addEventListener('click', e => {
      const b = e.target.closest('button'); if (!b) return;
      if (b.dataset.g === 'back') { closeGame(); return; }
      const i = gboxes.indexOf(openBox) + (b.dataset.g === 'next' ? 1 : -1);
      if (gboxes[i]) openGame(gboxes[i]);
    });
    applyGrid();
  }
  const _show = show;
  window.__mtOnHash = reveal;
  window.mtShow = show;
  tl.addEventListener('click', e => { const b = e.target.closest('.mt-tab'); if (b) show(b.dataset.tab, { top: b.dataset.tab === current ? 0 : undefined }); });
  tl.addEventListener('keydown', e => {
    const ids = TABS.map(t => t.id); let i = ids.indexOf(current);
    if (e.key === 'ArrowRight') i = (i + 1) % ids.length;
    else if (e.key === 'ArrowLeft') i = (i - 1 + ids.length) % ids.length;
    else if (e.key === 'Home') i = 0;
    else if (e.key === 'End') i = ids.length - 1;
    else return;
    e.preventDefault(); show(ids[i], { scroll: false }); tabBtns[ids[i]].focus();
  });
  bn.addEventListener('click', e => {
    const b = e.target.closest('.mt-bt'); if (!b) return;
    if (b === more) { sheet.hidden = false; more.setAttribute('aria-expanded', 'true'); const f = sheet.querySelector('.mt-sheet-btn'); if (f) f.focus(); return; }
    show(b.dataset.tab, { top: b.dataset.tab === current ? 0 : undefined });
  });
  sheet.addEventListener('click', e => {
    if (e.target === sheet || e.target.closest('.mt-sheet-close')) { closeSheet(); more.focus(); return; }
    const b = e.target.closest('.mt-sheet-btn'); if (b) show(b.dataset.tab, { top: 0 });
  });
  document.addEventListener('keydown', e => { if (e.key === 'Escape' && !sheet.hidden) { closeSheet(); more.focus(); } });
  window.addEventListener('popstate', () => { if (!location.hash) show('picks', { push: false }); else reveal(location.hash); });
  let start = 'picks';
  try { start = sessionStorage.getItem('mt-tab') || 'picks'; } catch (e) {}
  if (!(location.hash && reveal(location.hash))) {
    show(start, { push: false, scroll: false });
    try { history.replaceState(null, '', '#tab-' + current); } catch (e) {}
  }
}
'''

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_formatted_runs(paragraph, text, base_color=None, base_size=Pt(10), base_font="Calibri"):
    # Strip tooltip popups and HTML markup cleanly for Word docx
    text = re.sub(r'<span class="tip-text"[^>]*>.*?</span>', '', text, flags=re.DOTALL | re.IGNORECASE)
    sub_lines = re.split(r'<br\s*/?>', text, flags=re.IGNORECASE)
    for line_idx, sub_line in enumerate(sub_lines):
        if line_idx > 0:
            paragraph.add_run().add_break()
        tokens = re.split(r'(<a\s+href=[\'"][^\'"]+[\'"][^>]*>.*?</a>|\[.*?\]\(.*?\)|<img[^>]*>|<strong>.*?</strong>|<b>.*?</b>|\*\*.*?\*\*|\*.*?\*|\\\$(?=[A-Za-z\\{]).*?\\\$|\$(?=[A-Za-z\\{]).*?\$)', sub_line)
        for tok in tokens:
            if not tok:
                continue

            m_img = re.match(r'<img[^>]*src=[\'"](?:https?://[^\'"]*/)?([a-zA-Z0-9_-]+)\.png[\'"][^>]*>', tok, flags=re.IGNORECASE)
            if m_img:
                abbr = m_img.group(1).lower()
                logo_path = Path(f'E:/dev/projects/NFL_Dashboard/data/logos/{abbr}.png')
                if logo_path.exists():
                    try:
                        run_img = paragraph.add_run()
                        run_img.add_picture(str(logo_path), width=Pt(11), height=Pt(11))
                        paragraph.add_run(' ')
                    except Exception:
                        pass
                continue

            m_strong = re.match(r'<(?:strong|b)>(.*?)</(?:strong|b)>', tok, flags=re.DOTALL | re.IGNORECASE)
            if m_strong:
                content = m_strong.group(1)
                content = re.sub(r'</?(?:span|abbr|div|p|i|em|a)[^>]*>', '', content, flags=re.IGNORECASE)
                run = paragraph.add_run(content)
                run.font.name = base_font
                run.font.size = base_size
                if base_color:
                    run.font.color.rgb = base_color
                run.font.bold = True
                continue

            m_a = re.match(r'<a\s+href=[\'"]([^\'"]+)[\'"][^>]*>(.*?)</a>', tok, flags=re.DOTALL | re.IGNORECASE)
            if m_a:
                link_url, link_label = m_a.groups()
                clean_label = re.sub(r'</?[^>]+>', '', link_label)
                clean_label = re.sub(r'[\*_]', '', clean_label)
                if link_url.startswith('#'):
                    anchor_name = link_url[1:]
                    w_ns = nsdecls('w')
                    sz_val = int(base_size.pt * 2) if hasattr(base_size, 'pt') else 17
                    hl_xml = f'<w:hyperlink {w_ns} w:anchor="{anchor_name}"><w:r><w:rPr><w:color w:val="1E40AF"/><w:u w:val="single"/><w:rFonts w:ascii="{base_font}" w:hAnsi="{base_font}"/><w:sz w:val="{sz_val}"/><w:b/></w:rPr><w:t>{clean_label}</w:t></w:r></w:hyperlink>'
                    try:
                        paragraph._p.append(parse_xml(hl_xml))
                    except Exception:
                        run = paragraph.add_run(clean_label)
                        run.font.name = base_font
                        run.font.size = base_size
                        run.font.color.rgb = RGBColor(30, 64, 175)
                        run.font.underline = True
                else:
                    run = paragraph.add_run(clean_label)
                    run.font.name = base_font
                    run.font.size = base_size
                    run.font.color.rgb = RGBColor(30, 64, 175)
                    run.font.underline = True
                continue

            tok_clean = re.sub(r'</?(?:span|abbr|div|p|i|em|a)[^>]*>', '', tok, flags=re.IGNORECASE)
            
            m_link = re.match(r'\[(.*?)\]\((.*?)\)', tok)
            if m_link:
                link_label, link_url = m_link.groups()
                clean_label = re.sub(r'[\*_]', '', link_label)
                if link_url.startswith('#'):
                    anchor_name = link_url[1:]
                    w_ns = nsdecls('w')
                    sz_val = int(base_size.pt * 2) if hasattr(base_size, 'pt') else 17
                    hl_xml = f'<w:hyperlink {w_ns} w:anchor="{anchor_name}"><w:r><w:rPr><w:color w:val="1E40AF"/><w:u w:val="single"/><w:rFonts w:ascii="{base_font}" w:hAnsi="{base_font}"/><w:sz w:val="{sz_val}"/><w:b/></w:rPr><w:t>{clean_label}</w:t></w:r></w:hyperlink>'
                    try:
                        paragraph._p.append(parse_xml(hl_xml))
                    except Exception:
                        run = paragraph.add_run(clean_label)
                        run.font.name = base_font
                        run.font.size = base_size
                        run.font.color.rgb = RGBColor(30, 64, 175)
                        run.font.underline = True
                else:
                    run = paragraph.add_run(clean_label)
                    run.font.name = base_font
                    run.font.size = base_size
                    run.font.color.rgb = RGBColor(30, 64, 175)
                    run.font.underline = True
                continue

            if not tok_clean:
                continue

            run = paragraph.add_run()
            run.font.name = base_font
            run.font.size = base_size
            if base_color:
                run.font.color.rgb = base_color
            
            if tok_clean.startswith('**') and tok_clean.endswith('**'):
                run.text = tok_clean[2:-2]
                run.font.bold = True
            elif tok_clean.startswith('*') and tok_clean.endswith('*'):
                run.text = tok_clean[1:-1]
                run.font.italic = True
            elif (tok_clean.startswith(r'\$') and tok_clean.endswith(r'\$')) or (tok_clean.startswith('$') and tok_clean.endswith('$')):
                if tok_clean.startswith(r'\$') and tok_clean.endswith(r'\$'):
                    inner = tok_clean[2:-2]
                else:
                    inner = tok_clean[1:-1]
                inner = inner.replace(r'\rightarrow', ' → ').replace('->', ' → ')
                run.text = inner.strip()
                run.font.bold = True
            else:
                run.text = tok_clean.replace(r'\rightarrow', ' → ').replace('->', ' → ')

def generate_docx(md_path, docx_path):
    if docx is None:
        print(f"python-docx is not installed; skipping DOCX export: {docx_path}")
        return
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    PRIMARY = RGBColor(16, 44, 87)       # Navy #102C57
    SECONDARY = RGBColor(53, 95, 142)    # Slate Blue #355F8E
    DARK_TEXT = RGBColor(30, 41, 59)     # Slate 800
    MUTED_TEXT = RGBColor(100, 116, 139) # Slate 500

    lines = Path(md_path).read_text(encoding='utf-8').splitlines()
    in_table = False
    table_lines = []
    in_code_block = False
    code_lines = []
    bm_id = 100
    pending_bm_name = None

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        rows = []
        for tl in tbl_lines:
            if re.match(r'^\s*\|?\s*:?-+:?\s*\|', tl):
                continue
            parts = [c.strip() for c in tl.strip().strip('|').split('|')]
            if parts:
                rows.append(parts)
        if not rows:
            return
        num_cols = max(len(r) for r in rows)
        tbl = doc.add_table(rows=len(rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        if num_cols == 6:
            col_widths = [Inches(1.3), Inches(1.1), Inches(0.9), Inches(1.5), Inches(1.1), Inches(1.6)]
        elif num_cols == 7:
            col_widths = [Inches(1.0), Inches(0.9), Inches(1.1), Inches(1.1), Inches(0.9), Inches(1.0), Inches(1.5)]
        else:
            col_widths = [Inches(7.0 / num_cols)] * num_cols

        for r_idx, r_data in enumerate(rows):
            is_hdr = (r_idx == 0)
            row = tbl.rows[r_idx]
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            if is_hdr:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            for c_idx in range(num_cols):
                cell = row.cells[c_idx]
                if c_idx < len(col_widths):
                    cell.width = col_widths[c_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                val = r_data[c_idx] if c_idx < len(r_data) else ''
                
                if is_hdr:
                    set_cell_background(cell, '102C57')
                    add_formatted_runs(p, val, base_color=RGBColor(255, 255, 255), base_size=Pt(8.5))
                    for run in p.runs:
                        run.font.bold = True
                else:
                    bg = 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF'
                    set_cell_background(cell, bg)
                    add_formatted_runs(p, val, base_color=DARK_TEXT, base_size=Pt(8.5))
                    if 'Consensus' in val or '4U' in val:
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = PRIMARY
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def flush_code(c_lines):
        if not c_lines:
            return
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        txt = "\n".join(c_lines)
        run = p.add_run(txt)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = SECONDARY

    def flush_html_table(html_content):
        nonlocal bm_id
        raw_rows_matches = list(re.finditer(r'<tr([^>]*)>(.*?)</tr>', html_content, flags=re.DOTALL | re.IGNORECASE))
        if not raw_rows_matches:
            return
        rows = []
        row_bms = []
        for match in raw_rows_matches:
            tr_attrs = match.group(1)
            m_tr_id = re.search(r'id=[\'"]([a-zA-Z0-9_-]+)[\'"]', tr_attrs)
            row_bms.append(m_tr_id.group(1) if m_tr_id else None)
            raw_cells = re.findall(r'<(?:th|td)[^>]*>(.*?)</(?:th|td)>', match.group(2), flags=re.DOTALL | re.IGNORECASE)
            if raw_cells:
                rows.append([c.strip() for c in raw_cells])
        if not rows:
            return
        num_cols = max(len(r) for r in rows)
        tbl = doc.add_table(rows=len(rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        col_widths = [Inches(7.2 / num_cols)] * num_cols

        for r_idx, r_data in enumerate(rows):
            is_hdr = (r_idx == 0)
            row = tbl.rows[r_idx]
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            if is_hdr:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            row_bm = row_bms[r_idx] if r_idx < len(row_bms) else None

            for c_idx in range(num_cols):
                cell = row.cells[c_idx]
                if c_idx < len(col_widths):
                    cell.width = col_widths[c_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                val = r_data[c_idx] if c_idx < len(r_data) else ''

                if c_idx == 0 and row_bm:
                    w_ns = nsdecls('w')
                    p._p.append(parse_xml(f'<w:bookmarkStart {w_ns} w:id="{bm_id}" w:name="{row_bm}"/>'))
                    p._p.append(parse_xml(f'<w:bookmarkEnd {w_ns} w:id="{bm_id}"/>'))
                    bm_id += 1

                if is_hdr:
                    set_cell_background(cell, '102C57')
                    add_formatted_runs(p, val, base_color=RGBColor(255, 255, 255), base_size=Pt(7.5))
                    for run in p.runs:
                        run.font.bold = True
                else:
                    bg = 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF'
                    set_cell_background(cell, bg)
                    add_formatted_runs(p, val, base_color=DARK_TEXT, base_size=Pt(7.5))
                    if 'Consensus' in val or '4U' in val or 'Grade A' in val:
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = PRIMARY
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith('```'):
            if in_code_block:
                flush_code(code_lines)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            in_table = True
            i += 1
            continue
        elif in_table:
            flush_table(table_lines)
            table_lines = []
            in_table = False

        if '<table' in stripped:
            html_table_lines = [stripped]
            while i + 1 < len(lines) and '</table>' not in lines[i]:
                i += 1
                html_table_lines.append(lines[i])
            full_tbl_html = "\n".join(html_table_lines)
            flush_html_table(full_tbl_html)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if '<div class="quick-nav-header"' in stripped or '<div class="global-rollup-bar"' in stripped or '<div class="rollup-controls"' in stripped:
            div_depth = 1
            while i + 1 < len(lines) and div_depth > 0:
                i += 1
                div_depth += lines[i].count('<div') - lines[i].count('</div')
            i += 1
            continue

        if '<div class="quick-pick-line">' in stripped or stripped.startswith('<div class="quick-pick-line">'):
            content = re.sub(r'</?div[^>]*>', '', stripped).strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            add_formatted_runs(p, content, base_color=DARK_TEXT, base_size=Pt(9.5))
            i += 1
            continue

        if stripped.startswith('<details'):
            m_id = re.search(r'id=[\'"]([a-zA-Z0-9_-]+)[\'"]', stripped)
            if m_id:
                pending_bm_name = m_id.group(1)
            i += 1
            continue

        if stripped.startswith('<summary>') or stripped.startswith('<summary '):
            summary_txt = re.sub(r'</?summary[^>]*>', '', stripped).strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            if pending_bm_name:
                w_ns = nsdecls('w')
                p._p.append(parse_xml(f'<w:bookmarkStart {w_ns} w:id="{bm_id}" w:name="{pending_bm_name}"/>'))
                p._p.append(parse_xml(f'<w:bookmarkEnd {w_ns} w:id="{bm_id}"/>'))
                bm_id += 1
                pending_bm_name = None
            add_formatted_runs(p, summary_txt, base_color=SECONDARY, base_size=Pt(11))
            for run in p.runs:
                run.font.bold = True
            i += 1
            continue

        if stripped == '</details>':
            pending_bm_name = None
            i += 1
            continue

        if stripped.startswith('<style') or stripped.startswith('<div') or stripped.startswith('</div') or stripped.startswith('<button') or stripped.startswith('</button') or stripped.startswith('<span') or stripped.startswith('</span'):
            m_div_anc = re.search(r'<(?:div|a)\s+id=[\'"]([a-zA-Z0-9_-]+)[\'"]', stripped)
            if m_div_anc:
                bm_name = m_div_anc.group(1)
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                w_ns = nsdecls('w')
                p._p.append(parse_xml(f'<w:bookmarkStart {w_ns} w:id="{bm_id}" w:name="{bm_name}"/>'))
                p._p.append(parse_xml(f'<w:bookmarkEnd {w_ns} w:id="{bm_id}"/>'))
                bm_id += 1
            i += 1
            continue

        bm_names = []
        for m_bm in re.finditer(r'\{#([a-zA-Z0-9_-]+)\}', stripped):
            bm_names.append(m_bm.group(1))
        stripped = re.sub(r'\{#([a-zA-Z0-9_-]+)\}', '', stripped).strip()

        for m_a in re.finditer(r'<a\s+id=[\'"]([a-zA-Z0-9_-]+)[\'"]\s*>\s*</a>', stripped):
            bm_names.append(m_a.group(1))
        stripped = re.sub(r'<a\s+id=[\'"]([a-zA-Z0-9_-]+)[\'"]\s*>\s*</a>', '', stripped).strip()

        if re.match(r'^---+$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run("━" * 60)
            r.font.color.rgb = MUTED_TEXT
            i += 1
            continue

        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            w_ns = nsdecls('w')
            for bm_name in bm_names:
                p._p.append(parse_xml(f'<w:bookmarkStart {w_ns} w:id="{bm_id}" w:name="{bm_name}"/>'))
                p._p.append(parse_xml(f'<w:bookmarkEnd {w_ns} w:id="{bm_id}"/>'))
                bm_id += 1
            add_formatted_runs(p, stripped[2:], base_color=PRIMARY, base_size=Pt(18))
            for run in p.runs:
                run.font.bold = True
            i += 1
            continue

        if stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            w_ns = nsdecls('w')
            for bm_name in bm_names:
                p._p.append(parse_xml(f'<w:bookmarkStart {w_ns} w:id="{bm_id}" w:name="{bm_name}"/>'))
                p._p.append(parse_xml(f'<w:bookmarkEnd {w_ns} w:id="{bm_id}"/>'))
                bm_id += 1
            add_formatted_runs(p, stripped[3:], base_color=PRIMARY, base_size=Pt(13))
            for run in p.runs:
                run.font.bold = True
            i += 1
            continue

        if stripped.startswith('### ') or stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            w_ns = nsdecls('w')
            for bm_name in bm_names:
                p._p.append(parse_xml(f'<w:bookmarkStart {w_ns} w:id="{bm_id}" w:name="{bm_name}"/>'))
                p._p.append(parse_xml(f'<w:bookmarkEnd {w_ns} w:id="{bm_id}"/>'))
                bm_id += 1
            txt = stripped[4:] if stripped.startswith('### ') else stripped[5:]
            add_formatted_runs(p, txt, base_color=SECONDARY, base_size=Pt(11))
            for run in p.runs:
                run.font.bold = True
            i += 1
            continue

        if re.match(r'^\s{2,}\*\s+', line) or re.match(r'^\s{2,}-\s+', line):
            content = re.sub(r'^\s{2,}[\*-]\s+', '', line)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(2)
            r_bullet = p.add_run("▪ ")
            r_bullet.font.color.rgb = SECONDARY
            add_formatted_runs(p, content, base_color=DARK_TEXT, base_size=Pt(9.5))
            i += 1
            continue

        if stripped.startswith('* ') or stripped.startswith('- '):
            content = stripped[2:]
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_after = Pt(3)
            r_bullet = p.add_run("• ")
            r_bullet.font.color.rgb = PRIMARY
            add_formatted_runs(p, content, base_color=DARK_TEXT, base_size=Pt(9.5))
            i += 1
            continue

        m_num = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if m_num:
            num, content = m_num.groups()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_after = Pt(3)
            r_num = p.add_run(f"{num}. ")
            r_num.font.bold = True
            r_num.font.color.rgb = PRIMARY
            add_formatted_runs(p, content, base_color=DARK_TEXT, base_size=Pt(9.5))
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_formatted_runs(p, stripped, base_color=DARK_TEXT, base_size=Pt(9.5))
        i += 1

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code(code_lines)

    doc.save(docx_path)
    print(f"Successfully generated DOCX from markdown: {docx_path}")

def generate_html(md_path, html_path):
    lines = Path(md_path).read_text(encoding='utf-8').splitlines()
    body_html = []
    outside_html = []
    in_table = False
    table_lines = []
    in_code = False
    code_lines = []
    in_ul = False
    in_ol = False

    def md_inline(txt):
        def replace_link(match):
            label = match.group(1)
            url = match.group(2)
            label = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', label)
            label = re.sub(r'\*(.*?)\*', r'<em>\1</em>', label)
            cls_name = "back-to-top" if "Back to" in label or "⬆" in label else "table-link"
            target_attr = ' target="_blank" rel="noopener noreferrer"' if url.startswith(('http://', 'https://')) else ''
            return f'<a href="{url}" class="{cls_name}"{target_attr}>{label}</a>'

        txt = re.sub(r'\[(.*?)\]\((.*?)\)', replace_link, txt)
        txt = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', txt)
        txt = re.sub(r'\*(.*?)\*', r'<em>\1</em>', txt)
        txt = txt.replace(r'\rightarrow', '&rarr;').replace('->', '&rarr;')
        txt = re.sub(r'\\?\$(?=[A-Za-z\\{])(.*?)\\?\$', r'<em>\1</em>', txt)  # math only; leave money like $30 alone
        txt = re.sub(r'\bATD\b', 'Anytime TD', txt)
        return txt

    def close_lists():
        nonlocal in_ul, in_ol
        res = []
        if in_ul:
            res.append('</ul>')
            in_ul = False
        if in_ol:
            res.append('</ol>')
            in_ol = False
        return res

    for line in lines:
        s = line.strip()
        if s.startswith('```'):
            if in_code:
                body_html.append(f'<pre><code>{"".join(code_lines)}</code></pre>')
                code_lines = []
                in_code = False
            else:
                body_html.extend(close_lists())
                in_code = True
                code_lines = []
            continue

        if in_code:
            code_lines.append(line + "\n")
            continue

        if s.startswith('|') and s.endswith('|'):
            body_html.extend(close_lists())
            table_lines.append(s)
            in_table = True
            continue
        elif in_table:
            t_html = ['<div class="table-responsive"><table class="sortable-table">']
            is_first = True
            for tl in table_lines:
                if re.match(r'^\s*\|?\s*:?-+:?\s*\|', tl):
                    continue
                cells = [c.strip() for c in tl.strip().strip('|').split('|')]
                tag = 'th' if is_first else 'td'
                if is_first:
                    t_html.append('<thead><tr>')
                else:
                    t_html.append('<tr>')
                for c in cells:
                    fmt = md_inline(c)
                    if not is_first:
                        if 'Cross-Show Consensus' in c:
                            fmt = f'<span class="badge badge-cross">{fmt}</span>'
                        elif 'Show Consensus' in c or 'Consensus' in c or '4U' in c:
                            fmt = f'<span class="badge badge-consensus">{fmt}</span>'
                        elif '2 Units' in c or '2U' in c:
                            fmt = f'<span class="badge badge-unit">{fmt}</span>'
                        elif 'Teaser' in c:
                            fmt = f'<span class="badge badge-teaser">{fmt}</span>'
                        elif 'Best Bet' in c or 'Feature' in c or 'Simon Says' in c:
                            fmt = f'<span class="badge badge-best">{fmt}</span>'
                    t_html.append(f'<{tag}>{fmt}</{tag}>')
                if is_first:
                    t_html.append('</tr></thead><tbody>')
                    is_first = False
                else:
                    t_html.append('</tr>')
            t_html.append('</tbody></table></div>')
            body_html.append(''.join(t_html))
            table_lines = []
            in_table = False

        if not s:
            continue

        if s.startswith('<summary'):
            body_html.extend(close_lists())
            m_sum = re.match(r'<summary(?:\s+class=[\'"][^\'"]*[\'"])?>(.*?)</summary>', s)
            if m_sum:
                content = m_sum.group(1)
                body_html.append(f'<summary><span class="sum-text">{md_inline(content)}</span></summary>')
            else:
                body_html.append(s)
            continue

        if s.startswith('<div class="side-toc"') or s.startswith('<button class="toc-toggle"'):
            outside_html.append(s)
            continue

        if s.startswith('<style') or s.startswith('<details') or s.startswith('</details') or s.startswith('</summary') or s.startswith('<div') or s.startswith('</div') or s.startswith('<button') or s.startswith('</button') or s.startswith('<a id='):
            body_html.extend(close_lists())
            body_html.append(s)
            continue

        anchor_id = None
        m_anc = re.search(r'\{#([a-zA-Z0-9_-]+)\}', s)
        if m_anc:
            anchor_id = m_anc.group(1)
            s = re.sub(r'\{#([a-zA-Z0-9_-]+)\}', '', s).strip()
        m_a = re.search(r'<a\s+id=[\'"]([a-zA-Z0-9_-]+)[\'"]\s*>\s*</a>', s)
        if m_a:
            anchor_id = m_a.group(1)
            s = re.sub(r'<a\s+id=[\'"]([a-zA-Z0-9_-]+)[\'"]\s*>\s*</a>', '', s).strip()

        id_attr = f' id="{anchor_id}"' if anchor_id else ''

        if re.match(r'^---+$', s):
            body_html.extend(close_lists())
            body_html.append('<hr/>')
            continue

        if s.startswith('# '):
            body_html.extend(close_lists())
            body_html.append(f'<h1{id_attr}>{md_inline(s[2:])}</h1>')
            continue

        if s.startswith('## '):
            body_html.extend(close_lists())
            body_html.append(f'<h2{id_attr}>{md_inline(s[3:])}</h2>')
            continue

        if s.startswith('### '):
            body_html.extend(close_lists())
            body_html.append(f'<h3{id_attr}>{md_inline(s[4:])}</h3>')
            continue

        if s.startswith('#### '):
            body_html.extend(close_lists())
            body_html.append(f'<h4{id_attr}>{md_inline(s[5:])}</h4>')
            continue

        if re.match(r'^\s{2,}[\*-]\s+', line):
            c = re.sub(r'^\s{2,}[\*-]\s+', '', line)
            if not in_ul:
                body_html.append('<ul>')
                in_ul = True
            body_html.append(f'<li class="sub-bullet">{md_inline(c)}</li>')
            continue

        if s.startswith('* ') or s.startswith('- '):
            c = s[2:]
            if not in_ul:
                body_html.extend(close_lists())
                body_html.append('<ul>')
                in_ul = True
            body_html.append(f'<li>{md_inline(c)}</li>')
            continue

        m_num = re.match(r'^(\d+)\.\s+(.*)', s)
        if m_num:
            _, c = m_num.groups()
            if not in_ol:
                body_html.extend(close_lists())
                body_html.append('<ol>')
                in_ol = True
            body_html.append(f'<li>{md_inline(c)}</li>')
            continue

        body_html.extend(close_lists())
        body_html.append(f'<p>{md_inline(s)}</p>')

    body_html.extend(close_lists())

    doc_title = "NFL Week 1 Master Betting Intelligence Report"
    for l in lines:
        if l.startswith('# '):
            clean = re.sub(r'^[#\s\U00010000-\U0010ffff\u2600-\u27ff\*]+', '', l).strip()
            if clean:
                doc_title = clean
            break

    mw = re.search(r'Week\s+(\d+)', doc_title)
    week_str = mw.group(1) if mw else '5'

    full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{doc_title}</title>
<link rel="icon" href="data:image/svg+xml,<svg xmlns=%22http://www.w3.org/2000/svg%22 viewBox=%220 0 100 100%22><text y=%22.9em%22 font-size=%2290%22>🏈</text></svg>">
<style>
  /* F-mi-dark (2026-10-03): one dark on-screen theme, Option A "Tracker slate + teal".
     Matches the Live Tracker surfaces + the Platinum Rose teal accent. Light lives only in @media print. */
  :root {{
    color-scheme: dark;
    --primary: #5eead4;        /* interactive text: links, active labels */
    --primary-light: #99f6e4;
    --secondary: #cbd5e1;      /* h3 / h4 */
    --accent: #00d2be;         /* Platinum Rose teal: active fills, focus, open edges */
    --on-accent: #0f172a;      /* text on teal (9.3:1) */
    --on-primary: #0f172a;
    --bg: #0f172a;             /* page */
    --card: #1e293b;           /* boxes */
    --raised: #273449;         /* table heads, buttons, hover */
    --border: #334155;
    --text: #f8fafc;           /* headings, strong */
    --body: #cbd5e1;           /* paragraphs */
    --muted: #94a3b8;
    --highlight: #1e293b;
    --zebra: rgba(148, 163, 184, 0.06);
    --pos: #34d399;
    --neg: #f87171;            /* not #ef4444: that fails AA on cards (3.9:1) */
    --warn: #fcd34d;
    --info: #38bdf8;
    --star: #fbbf24;
    --tap: 40px;               /* min height for buttons / pills */
  }}
  html {{
    scroll-behavior: auto;
  }}
  :target {{
    scroll-margin-top: 24px;
  }}
  @keyframes target-pulse {{
    0% {{ background-color: rgba(0, 210, 190, 0.2); border-radius: 6px; padding: 2px 6px; }}
    100% {{ background-color: transparent; }}
  }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    background-color: var(--bg);
    color: var(--body);
    line-height: 1.65;
    margin: 0;
    padding: 32px 20px;
  }}
  .container {{
    max-width: 1100px;
    margin: 0 auto;
    background: var(--bg);
    border: 1px solid var(--border);
    border-radius: 14px;
    padding: 36px 44px;
    box-shadow: none;
  }}
  a {{ color: var(--primary); }}
  a:hover {{ color: var(--primary-light); }}
  a:focus-visible, button:focus-visible, summary:focus-visible {{ outline: 2px solid var(--accent); outline-offset: 2px; }}
  strong, b {{ color: var(--text); }}
  h1 {{
    color: var(--text);
    font-size: 26px;
    margin-top: 0;
    margin-bottom: 12px;
    line-height: 1.3;
  }}
  h2 {{
    color: var(--text);
    font-size: 20px;
    margin-top: 32px;
    margin-bottom: 16px;
    border-bottom: 2px solid var(--border);
    padding-bottom: 8px;
  }}
  h3 {{
    color: var(--secondary);
    font-size: 16px;
    margin-top: 20px;
    margin-bottom: 8px;
  }}
  h4 {{
    color: var(--secondary);
    font-size: 14.5px;
    margin-top: 16px;
    margin-bottom: 6px;
  }}
  p {{
    margin: 10px 0;
    font-size: 14.5px;
  }}
  hr {{
    border: 0;
    height: 1px;
    background: var(--border);
    margin: 28px 0;
  }}
  .table-responsive {{
    overflow-x: auto;
    overflow-y: visible;
    margin: 16px 0;
    padding-bottom: 24px;
  }}
  table {{
    width: 100%;
    border-collapse: collapse;
    font-size: 13.5px;
  }}
  th, td {{
    padding: 10px 12px;
    text-align: left;
    border-bottom: 1px solid var(--border);
  }}
  th {{
    background-color: var(--raised);
    color: var(--muted);
    font-size: 11.5px;
    letter-spacing: 0.05em;
    text-transform: uppercase;
    font-weight: 700;
    position: relative;
  }}
  th:hover, th:has(.term-tooltip:hover) {{
    z-index: 100;
  }}
  tr:nth-child(even) {{
    background-color: var(--zebra);
  }}
  .badge {{
    display: inline-block;
    padding: 3px 8px;
    border-radius: 4px;
    font-size: 11px;
    font-weight: 700;
    line-height: 1.2;
    margin: 1px 0;
  }}
  .badge-cross {{
    background: rgba(248, 113, 113, 0.14);
    color: #fca5a5;
    border: 1px solid rgba(248, 113, 113, 0.4);
  }}
  .badge-consensus {{
    background: rgba(56, 189, 248, 0.14);
    color: #7dd3fc;
    border: 1px solid rgba(56, 189, 248, 0.4);
  }}
  .badge-unit {{
    background: rgba(252, 211, 77, 0.14);
    color: #fcd34d;
    border: 1px solid rgba(252, 211, 77, 0.4);
  }}
  .badge-teaser {{
    background: rgba(192, 132, 252, 0.14);
    color: #d8b4fe;
    border: 1px solid rgba(192, 132, 252, 0.4);
  }}
  .badge-best {{
    background: rgba(52, 211, 153, 0.14);
    color: #6ee7b7;
    border: 1px solid rgba(52, 211, 153, 0.4);
  }}
  .badge-stable {{
    background: rgba(52, 211, 153, 0.14);
    color: #6ee7b7;
    border: 1px solid rgba(52, 211, 153, 0.4);
  }}
  .badge-watch {{
    background: rgba(252, 211, 77, 0.14);
    color: #fcd34d;
    border: 1px solid rgba(252, 211, 77, 0.4);
  }}
  .badge-volatile {{
    background: rgba(248, 113, 113, 0.14);
    color: #fca5a5;
    border: 1px solid rgba(248, 113, 113, 0.4);
  }}
  .badge-zero {{
    background: rgba(148, 163, 184, 0.14);
    color: #cbd5e1;
    border: 1px solid rgba(148, 163, 184, 0.4);
  }}
  .table-link {{
    color: var(--primary);
    text-decoration: none;
    font-weight: 600;
    border-bottom: 1.5px dashed var(--primary);
    transition: all 0.2s ease;
  }}
  .table-link:hover {{
    color: var(--primary-light);
    border-bottom-style: solid;
  }}
  .badge a {{
    color: inherit;
    text-decoration: none;
    border-bottom: none;
  }}
  .badge a:hover {{
    text-decoration: underline;
    border-bottom: none;
  }}
  .back-to-top {{
    display: inline-block;
    margin: 8px 0 16px 0;
    font-size: 12px;
    font-weight: 600;
    color: var(--primary);
    text-decoration: none;
    background: var(--highlight);
    border: 1px solid var(--border);
    padding: 8px 14px;
    border-radius: 8px;
    transition: all 0.2s ease;
  }}
  .back-to-top:hover {{
    background: var(--accent);
    color: var(--on-accent);
    border-color: var(--accent);
  }}
  ul, ol {{
    padding-left: 24px;
    margin: 10px 0;
  }}
  li {{
    margin-bottom: 8px;
    font-size: 14.5px;
  }}
  li.sub-bullet {{
    list-style-type: square;
    margin-bottom: 6px;
  }}
  pre {{
    background: var(--highlight);
    border: 1px solid var(--border);
    border-radius: 6px;
    padding: 12px 16px;
    font-size: 13px;
    overflow-x: auto;
  }}
  code {{
    font-family: Consolas, Monaco, monospace;
  }}
  /* Quick Nav Panel & Compact Controls */
  .quick-nav-panel {{
    background: var(--highlight);
    border: 1px solid var(--border);
    border-radius: 10px;
    padding: 12px 18px;
    margin: 18px 0 22px 0;
    box-shadow: 0 2px 4px rgba(0, 0, 0, 0.04);
  }}
  .quick-nav-header {{
    display: flex;
    flex-wrap: wrap;
    align-items: center;
    justify-content: space-between;
    gap: 10px;
    padding-bottom: 10px;
    margin-bottom: 10px;
    border-bottom: 1px solid var(--border);
  }}
  .quick-nav-controls {{
    display: flex;
    flex-wrap: wrap;
    align-items: center;
    gap: 8px;
  }}
  .quick-nav-label {{
    font-size: 13px;
    font-weight: 700;
    color: var(--primary);
    margin-right: 4px;
  }}
  .quick-nav-hint {{
    font-size: 12px;
    color: var(--muted);
  }}
  .btn-pill {{
    background: var(--raised);
    color: var(--text);
    border: 1px solid var(--border);
    min-height: var(--tap);
    padding: 0 16px;
    border-radius: 999px;
    font-size: 13.5px;
    font-weight: 700;
    cursor: pointer;
    transition: all 0.2s ease;
    display: inline-flex;
    align-items: center;
    gap: 4px;
  }}
  .btn-pill:hover {{
    border-color: var(--accent);
    color: var(--primary);
  }}
  .btn-pill.btn-primary {{
    background: var(--accent);
    color: var(--on-accent);
    border-color: var(--accent);
  }}
  .btn-pill.btn-primary:hover {{
    background: var(--primary);
    border-color: var(--primary);
    color: var(--on-accent);
  }}
  .quick-picks-list {{
    display: flex;
    flex-direction: column;
    gap: 8px;
  }}
  .quick-pick-line {{
    display: flex;
    flex-wrap: wrap;
    align-items: center;
    gap: 8px;
    font-size: 13.5px;
    line-height: 1.4;
  }}
  .pick-tier-badge {{
    font-size: 11px;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    padding: 2px 8px;
    border-radius: 4px;
    min-width: 72px;
    text-align: center;
  }}
  .badge-top5 {{
    background: rgba(52, 211, 153, 0.14);
    color: #6ee7b7;
    border: 1px solid rgba(52, 211, 153, 0.4);
  }}
  .badge-alt {{
    background: rgba(148, 163, 184, 0.14);
    color: #cbd5e1;
    border: 1px solid rgba(148, 163, 184, 0.4);
  }}
  .quick-pick-link {{
    display: inline-flex;
    align-items: center;
    gap: 4px;
    color: var(--primary);
    font-weight: 700;
    text-decoration: none;
    padding: 2px 6px;
    border-radius: 5px;
    transition: all 0.15s ease;
  }}
  .quick-pick-link:hover {{
    background: rgba(0, 210, 190, 0.12);
    color: var(--primary-light);
    transform: translateY(-1px);
  }}
  .quick-pick-sep {{
    color: var(--muted);
    font-weight: 400;
    user-select: auto;
  }}
  tr:target {{
    background-color: rgba(0, 210, 190, 0.16) !important;
    outline: 2px solid var(--accent);
    transition: background-color 0.5s ease;
  }}
  /* Rollup / Collapsible Box Styles */
  .global-rollup-bar {{
    display: flex;
    flex-wrap: wrap;
    align-items: center;
    gap: 10px;
    background: var(--highlight);
    border: 1px solid var(--border);
    border-radius: 8px;
    padding: 10px 16px;
    margin: 20px 0 24px 0;
  }}
  .global-rollup-bar .bar-title {{
    font-weight: 700;
    font-size: 13.5px;
    color: var(--primary);
  }}
  .global-rollup-bar .bar-hint {{
    font-size: 12.5px;
    color: var(--muted);
    margin-left: auto;
  }}
  .rollup-controls {{
    display: flex;
    gap: 8px;
    margin: 12px 0 16px 0;
  }}
  .btn-toggle {{
    background: var(--raised);
    color: var(--text);
    border: 1px solid var(--border);
    min-height: var(--tap);
    padding: 0 16px;
    border-radius: 10px;
    font-size: 13.5px;
    font-weight: 700;
    cursor: pointer;
    display: inline-flex;
    align-items: center;
    gap: 6px;
    transition: border-color 0.2s ease, color 0.2s ease, background 0.2s ease;
  }}
  .btn-toggle:hover {{
    border-color: var(--accent);
    color: var(--primary);
  }}
  .btn-primary, .btn-toggle.btn-primary {{
    background: var(--accent);
    color: var(--on-accent);
    border-color: var(--accent);
  }}
  .btn-primary:hover, .btn-toggle.btn-primary:hover {{
    background: var(--primary);
    border-color: var(--primary);
    color: var(--on-accent);
  }}
  .rollup-controls {{ flex-wrap: wrap; align-items: center; }}
  details.rollup-box {{
    background: var(--card);
    border: 1px solid var(--border);
    border-radius: 8px;
    margin: 14px 0;
    overflow: hidden;
    transition: border-color 0.2s ease, box-shadow 0.2s ease;
  }}
  details.rollup-box:hover {{
    border-color: var(--primary-light);
  }}
  details.rollup-box[open] {{
    border-color: rgba(0, 210, 190, 0.5);
    box-shadow: inset 3px 0 0 var(--accent);
  }}
  details.rollup-box:target {{
    border-color: var(--accent) !important;
    box-shadow: 0 0 0 4px rgba(0, 210, 190, 0.3) !important;
    animation: target-box-pulse 2.5s ease-out;
  }}
  @keyframes target-box-pulse {{
    0% {{ box-shadow: 0 0 0 6px rgba(0, 210, 190, 0.5); border-color: var(--accent); }}
    100% {{ box-shadow: inset 3px 0 0 var(--accent); }}
  }}
  details.rollup-box summary {{
    min-height: 52px;
    box-sizing: border-box;
    gap: 12px;
    padding: 10px 14px 10px 18px;
    font-size: 15.5px;
    font-weight: 700;
    color: var(--text);
    background: var(--card);
    cursor: pointer;
    user-select: text;
    display: flex;
    align-items: center;
    justify-content: space-between;
    list-style: none;
    transition: background 0.2s ease;
  }}
  details.rollup-box summary::-webkit-details-marker {{
    display: none;
  }}
  details.rollup-box summary:hover {{
    background: var(--raised);
  }}
  details.rollup-box > summary .sum-text {{ flex: 1 1 auto; min-width: 0; }}
  details.rollup-box > summary::after {{
    content: '▾';
    flex: 0 0 30px;
    width: 30px;
    height: 30px;
    display: inline-grid;
    place-items: center;
    border-radius: 8px;
    background: var(--raised);
    border: 1px solid var(--border);
    font-size: 15px;
    color: var(--accent);
    transition: transform 0.2s ease;
    transform: rotate(-90deg);
  }}
  details.rollup-box[open] > summary::after {{
    transform: rotate(0deg);
  }}
  details.rollup-box .rollup-content {{
    padding: 16px 20px;
    border-top: 1px solid var(--border);
  }}
  /* Interactive Laymen Term Tooltips */
  .term-tooltip {{
    position: relative;
    display: inline-flex;
    align-items: baseline;
    gap: 3px;
    cursor: help;
    border-bottom: 1.5px dotted var(--info);
    color: inherit;
    font-weight: inherit;
  }}
  .term-tooltip:hover {{
    z-index: 101;
  }}
  .term-tooltip::after {{
    content: ' ⓘ';
    font-size: 10px;
    color: var(--info);
    opacity: 0.85;
    font-weight: bold;
  }}
  .term-tooltip .tip-text {{
    visibility: hidden;
    width: min(280px, 85vw);
    max-width: 85vw;
    box-sizing: border-box;
    background-color: #0b1222;
    color: #f8fafc;
    text-align: left;
    border-radius: 8px;
    padding: 10px 14px;
    position: absolute;
    z-index: 99999;
    top: calc(100% + 8px);
    bottom: auto;
    left: 50%;
    transform: translateX(-50%);
    opacity: 0;
    transition: opacity 0.2s ease, transform 0.2s ease;
    font-size: 12px;
    font-weight: 400;
    line-height: 1.45;
    box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.5), 0 8px 10px -6px rgba(0, 0, 0, 0.3);
    border: 1px solid #334155;
    pointer-events: none;
    white-space: normal;
    text-transform: none;
  }}
  .term-tooltip .tip-text strong {{
    color: var(--primary);
    display: block;
    margin-bottom: 4px;
    font-size: 12.5px;
    border-bottom: 1px solid #1e293b;
    padding-bottom: 3px;
  }}
  .term-tooltip .tip-text::after {{
    content: "";
    position: absolute;
    bottom: 100%;
    left: 50%;
    margin-left: -6px;
    border-width: 6px;
    border-style: solid;
    border-color: transparent transparent #0b1222 transparent;
  }}
  .term-tooltip:hover .tip-text {{
    visibility: visible;
    opacity: 1;
    transform: translateX(-50%) translateY(3px);
  }}
  th:first-child .term-tooltip .tip-text,
  th:nth-child(2) .term-tooltip .tip-text {{
    left: 0;
    transform: none;
  }}
  th:first-child .term-tooltip .tip-text::after,
  th:nth-child(2) .term-tooltip .tip-text::after {{
    left: 24px;
  }}
  th:first-child .term-tooltip:hover .tip-text,
  th:nth-child(2) .term-tooltip:hover .tip-text {{
    transform: translateY(3px);
  }}
  th:last-child .term-tooltip .tip-text,
  th:nth-last-child(2) .term-tooltip .tip-text {{
    left: auto;
    right: 0;
    transform: none;
  }}
  th:last-child .term-tooltip .tip-text::after,
  th:nth-last-child(2) .term-tooltip .tip-text::after {{
    left: auto;
    right: 24px;
  }}
  th:last-child .term-tooltip:hover .tip-text,
  th:nth-last-child(2) .term-tooltip:hover .tip-text {{
    transform: translateY(3px);
  }}
  details.section-alternates {{
    margin-top: 18px;
    border: 1.5px dashed #475569;
    border-radius: 8px;
  }}
  details.section-alternates summary {{
    background: rgba(30, 41, 59, 0.7);
    color: var(--primary);
    font-size: 14.5px;
  }}
  /* Sortable Tables */
  th.sortable-header {{
    cursor: pointer;
    user-select: text;
    position: relative;
    padding-right: 22px;
    transition: background-color 0.2s ease;
  }}
  th.sortable-header:hover {{
    background-color: var(--border);
    color: var(--text);
  }}
  .sort-icon {{
    display: inline-block;
    margin-left: 6px;
    font-size: 11px;
    opacity: 0.65;
    vertical-align: middle;
    transition: opacity 0.2s ease, color 0.2s ease;
  }}
  th.sortable-header:hover .sort-icon {{
    opacity: 1;
  }}
  th.sortable-header[data-sort-dir="asc"] .sort-icon,
  th.sortable-header[data-sort-dir="desc"] .sort-icon {{
    opacity: 1;
    color: var(--accent);
    font-weight: 900;
  }}
  .team-logo {{
    width: 22px;
    height: 22px;
    vertical-align: middle;
    margin-right: 6px;
    object-fit: contain;
    display: inline-block;
    filter: drop-shadow(0 1px 2px rgba(0, 0, 0, 0.25));
  }}
  .team-logo-sm {{
    width: 18px;
    height: 18px;
    vertical-align: middle;
    margin-right: 4px;
    object-fit: contain;
    display: inline-block;
    filter: drop-shadow(0 1px 1px rgba(0, 0, 0, 0.2));
  }}
  /* Target pulse highlight for Jump to a Game / anchor links */
  @keyframes target-pulse {{
    0% {{ outline: 3px solid var(--accent); box-shadow: 0 0 0 6px rgba(0, 210, 190, 0.4); }}
    70% {{ outline: 3px solid var(--accent); box-shadow: 0 0 0 6px rgba(0, 210, 190, 0.2); }}
    100% {{ outline: none; box-shadow: none; }}
  }}
  .target-pulse {{
    animation: target-pulse 2s ease-out;
  }}

  /* Matchup Popover Modal Dashboard */
  .matchup-modal-overlay {{
    position: fixed;
    top: 0; left: 0; right: 0; bottom: 0;
    background: rgba(11, 17, 32, 0.88);
    backdrop-filter: blur(10px);
    -webkit-backdrop-filter: blur(10px);
    z-index: 10000;
    display: none;
    align-items: center;
    justify-content: center;
    padding: 18px;
    box-sizing: border-box;
  }}
  .matchup-modal-overlay.active {{
    display: flex;
  }}
  .matchup-modal-dialog {{
    background: var(--card);
    border: 1px solid var(--border);
    border-radius: 14px;
    width: 100%;
    max-width: 920px;
    max-height: 88vh;
    display: flex;
    flex-direction: column;
    box-shadow: 0 24px 60px rgba(0, 0, 0, 0.65);
    overflow: hidden;
    animation: modal-pop 0.2s cubic-bezier(0.16, 1, 0.3, 1);
  }}
  @keyframes modal-pop {{
    0% {{ opacity: 0; transform: scale(0.96) translateY(10px); }}
    100% {{ opacity: 1; transform: scale(1) translateY(0); }}
  }}
  .matchup-modal-header {{
    padding: 14px 18px;
    border-bottom: 1px solid var(--border);
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 12px;
    background: var(--raised);
    flex-shrink: 0;
  }}
  .matchup-modal-title {{
    font-size: 15.5px;
    font-weight: 750;
    color: var(--text);
    margin: 0;
    display: flex;
    align-items: center;
    gap: 8px;
    flex: 1;
    min-width: 0;
    overflow: hidden;
    text-overflow: ellipsis;
    white-space: nowrap;
  }}
  .matchup-modal-actions {{
    display: flex;
    align-items: center;
    gap: 8px;
    flex-shrink: 0;
  }}
  .matchup-modal-body {{
    padding: 20px;
    overflow-y: auto;
    flex: 1;
    font-size: 14px;
    line-height: 1.55;
    color: var(--body);
  }}
  .matchup-modal-body details.rollup-box {{
    margin: 10px 0;
  }}
  .matchup-modal-body summary {{
    min-height: 44px;
    font-size: 14px;
  }}

  /* Confidence Stars Visual Hierarchy */
  .star-rating {{
    display: inline-flex;
    align-items: center;
    gap: 4px;
    font-feature-settings: "tnum";
    white-space: nowrap;
    padding: 1px 6px;
    border-radius: 4px;
    transition: all 0.2s ease;
  }}
  .star-tier-high {{
    font-weight: 850;
    color: #fbbf24;
    background: rgba(245, 158, 11, 0.12);
    border: 1px solid rgba(245, 158, 11, 0.35);
    box-shadow: 0 0 10px rgba(245, 158, 11, 0.2);
  }}
  .star-tier-high .stars-active {{
    color: #f59e0b;
    font-size: 1.15em;
    letter-spacing: 1px;
    filter: drop-shadow(0 0 4px rgba(245, 158, 11, 0.65));
  }}
  .star-tier-mid {{
    font-weight: 700;
    color: var(--accent);
    background: rgba(0, 210, 190, 0.08);
    border: 1px solid rgba(0, 210, 190, 0.25);
  }}
  .star-tier-mid .stars-active {{
    color: var(--accent);
    letter-spacing: 0.5px;
  }}
  .star-tier-low {{
    font-weight: 500;
    color: #94a3b8;
    background: rgba(148, 163, 184, 0.06);
    border: 1px solid rgba(148, 163, 184, 0.18);
  }}
  .star-tier-low .stars-active {{
    color: #64748b;
  }}
  .stars-muted {{
    color: #334155;
    opacity: 0.55;
  }}

  /* At-a-Glance Picks Strip */
  .at-a-glance-bar {{
    display: flex;
    flex-direction: column;
    gap: 10px;
    margin: 16px 0 24px 0;
    padding: 14px 18px;
    background: var(--card);
    border: 1px solid var(--border);
    border-left: 4px solid var(--accent);
    border-radius: 12px;
  }}
  .at-a-glance-header {{
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 10px;
    flex-wrap: wrap;
  }}
  .at-a-glance-title {{
    font-size: 12px;
    font-weight: 800;
    letter-spacing: 0.08em;
    text-transform: uppercase;
    color: var(--accent);
  }}
  .at-a-glance-pills {{
    display: flex;
    flex-wrap: wrap;
    gap: 8px;
  }}
  .at-glance-pill {{
    display: inline-flex;
    align-items: center;
    gap: 7px;
    padding: 6px 14px;
    background: var(--raised);
    border: 1px solid var(--border);
    border-radius: 999px;
    color: var(--text);
    font-size: 13px;
    font-weight: 750;
    text-decoration: none;
    transition: all 0.15s ease;
  }}
  .at-glance-pill:hover {{
    border-color: var(--accent);
    background: color-mix(in srgb, var(--accent) 15%, var(--raised));
    transform: translateY(-1px);
    color: var(--primary);
  }}
  .at-glance-pill .team-logo {{
    margin-right: 0;
    width: 20px;
    height: 20px;
  }}

  /* Card-Builder & Bookmarks Layer */
  .bm-btn {{
    appearance: none;
    display: inline-flex;
    align-items: center;
    justify-content: center;
    gap: 4px;
    min-height: 24px;
    padding: 3px 9px;
    border: 1.5px solid var(--accent);
    background: rgba(0, 210, 190, 0.12);
    color: var(--accent);
    border-radius: 999px;
    font-size: 11.5px;
    font-weight: 800;
    text-transform: uppercase;
    letter-spacing: 0.04em;
    cursor: pointer;
    white-space: nowrap;
    vertical-align: middle;
    line-height: 1;
    transition: all 0.15s ease;
    user-select: none;
    -webkit-user-select: none;
  }}
  .bm-btn:hover {{
    background: var(--accent);
    color: var(--on-accent);
    border-color: var(--accent);
    transform: translateY(-1px);
  }}
  .bm-btn.on {{
    background: #facc15;
    color: #0f172a;
    border-color: #facc15;
    font-weight: 850;
    box-shadow: 0 0 10px rgba(250, 204, 21, 0.35);
  }}
  .bm-btn.on:hover {{
    background: #fde047;
    border-color: #fde047;
  }}
  .bm-btn:focus-visible {{ outline: 2px solid var(--accent); outline-offset: 2px; }}
  .bm-btn .bm-icon {{ font-size: 13px; line-height: 1; }}
  .bm-btn .bm-text {{ display: none; }}
  .bm-btn {{ min-height: 24px; width: 26px; padding: 0; margin-right: 6px; border-width: 1px; border-radius: 6px; }}
  .bm-btn:hover {{ transform: none; }}
  .bm-btn .bm-icon {{ font-size: 14px; }}
  .at-glance-item .pill-bm-btn, .leg-bm-btn {{ width: 24px; min-height: 24px; padding: 0; }}
  .at-glance-item {{ display: inline-flex; align-items: center; gap: 4px; background: var(--card); border: 1px solid var(--border); border-radius: 999px; padding: 2px 6px 2px 2px; }}
  .at-glance-item .at-glance-pill {{ border: none; background: transparent; padding: 4px 6px 4px 8px; }}
  .at-glance-item .pill-bm-btn {{ padding: 3px 8px; min-height: 22px; font-size: 10.5px; }}
  .sc-tile-wrap {{ display: flex; flex-direction: column; gap: 5px; }}
  .sc-tile-wrap .sc-tile {{ flex: 1; }}
  .sc-tile-btn {{ width: 100%; justify-content: center; padding: 6px 8px; font-size: 11.5px; border-radius: 8px; }}
  .pk-quick {{ display: flex; flex-direction: column; gap: 8px; padding: 12px 14px; margin: 8px 0 14px; border: 1px solid var(--accent); border-radius: 12px; background: var(--card); }}
  .pk-qrow {{ display: flex; flex-wrap: wrap; align-items: center; gap: 8px; }}
  .pk-qlabel {{ flex: 0 0 110px; font-weight: 700; font-size: .85em; color: var(--muted); }}
  .pk-chip {{ display: inline-flex; align-items: center; gap: 6px; min-height: 34px; padding: 4px 12px; border: 1px solid var(--border); border-radius: 999px; background: var(--raised); color: var(--text); text-decoration: none; }}
  .pk-chip:hover {{ border-color: var(--accent); }}
  .pk-chip b {{ white-space: nowrap; }}
  .sc-group.rollup-box > summary {{ font-size: 1.05em; }}
  .sc-topbar {{ position: sticky; top: 0; z-index: 60; display: flex; align-items: center; gap: 8px; padding: 8px 14px; background: var(--card); border-bottom: 2px solid var(--accent); overflow-x: auto; white-space: nowrap; }}
  body.sc-tabs-on {{ padding-top: 0; }}
  body.sc-tabs-on .sc-topbar {{ margin: 0 0 0 -20px; }}
  body.sc-tabs-on .container {{ margin-top: 24px; }}
  .sc-who {{ display:inline-flex; align-items:center; gap:4px; margin-right:10px; padding:3px 4px 3px 10px; border:1px solid var(--border); border-radius:999px; background:var(--card); flex:none; }}
  .sc-who-l {{ font-size:.72rem; color:var(--muted); margin-right:2px; }}
  .sc-who button {{ border:0; background:transparent; color:var(--text); font:inherit; font-size:.8rem; font-weight:700; padding:4px 11px; border-radius:999px; cursor:pointer; }}
  .sc-who button.on {{ background:var(--accent); color:#06231f; }}
  .sc-who.flash {{ animation:scwho 0.5s 3; box-shadow:0 0 0 2px var(--accent); }}
  @keyframes scwho {{ 50% {{ box-shadow:0 0 0 6px rgba(0,210,190,.35); }} }}
  body[data-who="andy"] input.pk-box[data-who="andy"], body[data-who="amanda"] input.pk-box[data-who="amanda"] {{ outline:2px solid var(--accent); outline-offset:1px; }}
  .sc-cb {{ display:inline-flex; flex-direction:column; gap:3px; margin-right:8px; vertical-align:middle; }}
  .sc-cb label {{ display:flex; align-items:center; gap:5px; font-size:.72rem; font-weight:700; color:var(--muted); cursor:pointer; white-space:nowrap; }}
  .sc-cb label.on {{ color:var(--text); }}
  .sc-cb input {{ width:16px; height:16px; accent-color:var(--accent); cursor:pointer; margin:0; }}
  .sc-pop {{ position:absolute; z-index:60; display:flex; gap:6px; padding:6px; border:1px solid var(--border); border-radius:10px; background:var(--card); box-shadow:0 8px 24px rgba(0,0,0,.45); }}
  .sc-pop button {{ border:1px solid var(--border); background:transparent; color:var(--text); border-radius:8px; padding:6px 10px; font:inherit; font-size:.8rem; font-weight:700; cursor:pointer; }}
  .sc-pop button.on {{ background:var(--accent); color:#06231f; border-color:var(--accent); }}
  .sp-legend {{ display:flex; gap:16px; margin:6px 0 10px; font-size:.78rem; color:var(--muted); }}
  .sp-legend span {{ display:inline-flex; align-items:center; gap:6px; }}
  .sp-k {{ width:12px; height:12px; border-radius:3px; display:inline-block; }}
  .sp-k.bets, .sp-row.bets .sp-bar i {{ background:#38bdf8; }}
  .sp-k.money, .sp-row.money .sp-bar i {{ background:#34d399; }}
  .sp-grid {{ display:grid; gap:10px; }}
  .sp-game {{ border:1px solid var(--border); border-radius:10px; background:var(--card); padding:10px 12px; }}
  .sp-head {{ display:flex; justify-content:space-between; align-items:baseline; gap:8px; margin-bottom:8px; font-size:.95rem; }}
  .sp-head span {{ font-size:.75rem; color:var(--muted); }}
  .sp-cols {{ display:grid; grid-template-columns:repeat(auto-fit, minmax(170px, 1fr)); gap:14px; }}
  .sp-mt {{ font-size:.68rem; letter-spacing:.06em; text-transform:uppercase; color:var(--muted); margin-bottom:4px; }}
  .sp-side {{ margin-bottom:8px; }}
  .sp-name {{ font-size:.82rem; font-weight:600; }}
  .sp-name em {{ font-style:normal; font-weight:400; color:var(--muted); margin-left:4px; }}
  .sp-row {{ display:grid; grid-template-columns:38px 1fr 38px; align-items:center; gap:6px; font-size:.72rem; color:var(--muted); margin-top:3px; }}
  .sp-row b {{ color:var(--text); text-align:right; font-variant-numeric:tabular-nums; }}
  .sp-bar {{ height:9px; border-radius:5px; background:rgba(148,163,184,.18); overflow:hidden; display:block; }}
  .sp-bar i {{ display:block; height:100%; border-radius:5px; }}
  .sp-na {{ font-size:.78rem; color:var(--muted); padding:8px 0; }}
  body.sc-tabs-on .rollup-controls, body.sc-tabs-on .global-rollup-bar {{ display: none; }}
  .sc-tools {{ margin-left: auto; display: inline-flex; gap: 2px; flex: 0 0 auto; }}
  .sc-tool {{ width: 30px; height: 30px; padding: 0; border: 0; border-radius: 6px; background: transparent; color: var(--muted); font-size: 18px; line-height: 1; cursor: pointer; }}
  .sc-tool:hover {{ color: var(--text); background: var(--raised); }}
  .sc-brand {{ font-weight: 800; color: var(--text); margin-right: 8px; }}
  .sc-tabwrap {{ display: inline-flex; align-items: stretch; border-radius: 8px; background: var(--raised); flex: 0 0 auto; }}
  .sc-tabwrap .sc-tab {{ flex: 0 0 auto; border-radius: 8px 0 0 8px; background: transparent; padding: 8px 10px 8px 12px; }}
  .sc-tabwrap .sc-caret {{ border: 0; border-left: 1px solid var(--border); background: transparent; color: var(--body); padding: 0 10px; min-height: var(--tap); border-radius: 0 8px 8px 0; cursor: pointer; font: inherit; }}
  .sc-tabwrap .sc-caret:hover {{ color: var(--text); background: rgba(255,255,255,.06); }}
  .sc-tabwrap.active {{ background: var(--accent); }}
  .sc-tabwrap.active .sc-tab, .sc-tabwrap.active .sc-caret {{ color: var(--on-accent); }}
  .sc-tabwrap.active .sc-caret {{ border-left-color: rgba(15,23,42,.25); }}
  .sc-menu {{ position: fixed; z-index: 70; overflow-y: auto; background: var(--card); border: 1px solid var(--accent); border-radius: 10px; box-shadow: 0 12px 32px rgba(0,0,0,.5); padding: 6px; }}
  .sc-menu[hidden] {{ display: none; }}
  .sc-mi {{ display: block; padding: 8px 10px; border-radius: 6px; color: var(--text); text-decoration: none; font-size: 13.5px; line-height: 1.3; }}
  .sc-mi:hover {{ background: var(--raised); }}
  .sc-mi.sub {{ padding-left: 24px; color: var(--body); font-size: 12.5px; }}
  @media (max-width: 1599px) {{ .sc-brand {{ display: none; }} }}
  .sc-tabs {{ display: flex; flex-wrap: nowrap; overflow-x: auto; gap: 6px; margin: 14px 0 12px; padding: 6px; border: 1px solid var(--border); border-radius: 12px; background: var(--card); position: sticky; top: 0; z-index: 30; }}
  .sc-tab {{ flex: 1 0 auto; white-space: nowrap; min-height: var(--tap); padding: 8px 10px; font-size: 13px; border: 1px solid transparent; border-radius: 8px; background: var(--raised); color: var(--body); font: inherit; font-weight: 600; cursor: pointer; }}
  .sc-tab:hover {{ color: var(--text); }}
  .sc-tab.active {{ background: var(--accent); color: var(--on-accent); }}
  .sc-panel {{ display: none; }}
  .sc-panel.active {{ display: block; }}
  .pk-panel {{ position: fixed; top: 0; right: 0; bottom: 0; width: 450px; z-index: 40; display: flex; flex-direction: column; background: var(--card); border-left: 2px solid var(--accent); box-shadow: -8px 0 24px rgba(0,0,0,.35); }}
  .pk-ph {{ display: flex; align-items: center; gap: 8px; padding: 12px 14px; font-weight: 700; border-bottom: 1px solid var(--border); cursor: pointer; }}
  .pk-ph .pk-counts {{ margin-left: auto; font-weight: 500; font-size: .8em; color: var(--muted); }}
  .pk-caret {{ transition: transform .2s; }}
  .pk-body {{ display: flex; flex-direction: column; flex: 1 1 auto; min-height: 0; padding: 8px 12px 12px; gap: 8px; }}
  .pk-top {{ display: flex; justify-content: space-between; font-size: .9em; }}
  .pk-top b {{ color: var(--primary); }}
  .pk-statusline {{ font-size: .82em; color: var(--muted); min-height: 1.2em; }}
  .pk-btns {{ display: flex; flex-wrap: wrap; gap: 6px; }}
  .pk-btns .btn-toggle {{ padding: 5px 9px; font-size: 11.5px; min-height: 0; }}
  .pk-share {{ margin-top: 6px; border: 1px solid var(--border); border-radius: 8px; padding: 6px 8px; }}
  .pk-share summary {{ cursor: pointer; font-weight: 600; font-size: .85em; }}
  .pk-hint {{ font-size: .78em; color: var(--muted); margin: 6px 0; }}
  #pk-text {{ width: 100%; box-sizing: border-box; background: var(--bg); color: var(--text); border: 1px solid var(--border); border-radius: 6px; padding: 6px; font: inherit; font-size: 12px; margin-bottom: 6px; }}
  .pk-body {{ overflow-y: auto; }}
  .pk-scroll {{ flex: 1 1 auto; min-height: 120px; overflow-y: auto; border: 1px solid var(--border); border-radius: 8px; }}
  .pk-table {{ width: 100%; border-collapse: collapse; font-size: 12.5px; margin: 0; }}
  .pk-table th {{ position: sticky; top: 0; background: var(--raised); padding: 6px 6px; text-align: left; font-size: 11px; z-index: 1; }}
  .pk-table td {{ padding: 5px 6px; border-top: 1px solid var(--border); vertical-align: middle; }}
  .pk-table tr.pk-first td {{ border-top: 2px solid var(--border); }}
  .pk-table .pk-side {{ white-space: nowrap; width: 1%; }}
  .pk-table .pk-side .tl {{ margin-right: 5px; vertical-align: middle; }}
  .pk-table .pk-mk {{ color: var(--muted); font-size: 10.5px; }}
  .pk-table .pk-pair {{ white-space: nowrap; }}
  .pk-table .pk-pair label {{ display: inline-flex; align-items: center; gap: 2px; margin-right: 3px; font-size: 10.5px; font-weight: 700; cursor: pointer; }}
  .pk-table .pk-pair input {{ width: 13px; height: 13px; }}
  .pk-table td.pk-pair {{ padding-left: 3px; padding-right: 3px; }}
  .pk-table .pk-st {{ font-weight: 700; color: var(--primary); }}
  .pk-table .pk-pair input {{ margin: 0; }}
  .sc-slot {{ display: inline-block; vertical-align: middle; }}
  .sc-slot .sc-cb {{ margin-right: 6px; }}
  .pk-table .pk-st {{ font-size: 11px; white-space: nowrap; }}
  .pk-official {{ border-top: 1px solid var(--border); padding-top: 8px; }}
  .pk-oh {{ font-weight: 700; font-size: .9em; }}
  .pk-official ol {{ margin: 4px 0 6px 20px; padding: 0; font-size: .88em; }}
  .pk-empty {{ opacity: .55; list-style: none; margin-left: -18px; }}
  body.sc-tabs-on {{ padding-right: 462px; }}
  body.sc-tabs-on.pk-collapsed {{ padding-right: 0; }}
  body.sc-tabs-on .card-builder-trigger {{ display: none; }}
  .pk-panel.collapsed {{ top: auto; height: auto; width: 100%; border-left: 0; border-top: 2px solid var(--accent); }}
  .pk-panel.collapsed .pk-body {{ display: none; }}
  .pk-panel.collapsed .pk-caret {{ transform: rotate(180deg); }}
  @media (max-width: 1199px) {{
    body.sc-tabs-on {{ padding-right: 0; padding-bottom: 54px; }}
    .pk-panel {{ top: auto; height: 78vh; width: 100%; border-left: 0; border-top: 2px solid var(--accent); }}
    .pk-panel.collapsed {{ height: auto; }}
  }}
  .sc-pick-btn {{ margin-left: 8px; padding: 3px 10px; font-size: 11.5px; }}
  .leg-bm-btn {{ padding: 3px 8px; min-height: 22px; font-size: 10.5px; }}
  .bm-flash {{ animation: bmflash 2.4s ease-out 1; }}
  @keyframes bmflash {{ 0% {{ background: rgba(250,204,21,.35); }} 100% {{ background: transparent; }} }}
  .bm-count {{
    display: inline-block;
    min-width: 18px;
    margin-left: 5px;
    padding: 0 6px;
    border-radius: 999px;
    background: #facc15;
    color: #0f172a;
    font-size: 11px;
    font-weight: 800;
    line-height: 18px;
    text-align: center;
  }}
  .bm-count:empty {{ display: none; }}
  .bm-group {{ margin: 16px 0; }}
  .bm-group h3 {{ margin: 0 0 8px; font-size: 14.5px; color: var(--text); border-bottom: 1px solid var(--border); padding-bottom: 4px; }}
  .bm-row {{
    display: flex;
    gap: 10px;
    align-items: flex-start;
    padding: 10px 12px;
    border: 1px solid var(--border);
    border-radius: 9px;
    background: var(--raised);
    margin: 6px 0;
  }}
  .bm-row a {{
    color: var(--text);
    text-decoration: none;
    flex: 1;
    min-width: 0;
    overflow-wrap: anywhere;
    font-size: 13.5px;
  }}
  .bm-row a:hover {{ color: var(--primary); }}
  .bm-row .bm-ctx {{ display: block; color: var(--muted); font-size: 12px; margin-top: 3px; }}

  /* Card Builder Slide-Out Drawer */
  .card-builder-trigger {{
    position: fixed;
    bottom: 24px;
    right: 24px;
    z-index: 9998;
    display: inline-flex;
    align-items: center;
    gap: 8px;
    padding: 11px 20px;
    background: var(--accent);
    color: var(--on-accent);
    border: none;
    border-radius: 999px;
    font-size: 13.5px;
    font-weight: 800;
    box-shadow: 0 8px 24px rgba(0, 0, 0, 0.55);
    cursor: pointer;
    transition: transform 0.15s ease, box-shadow 0.15s ease;
  }}
  .card-builder-trigger:hover {{
    transform: translateY(-2px);
    box-shadow: 0 12px 32px rgba(0, 0, 0, 0.75);
  }}
  .card-drawer-overlay {{
    position: fixed;
    top: 0; left: 0; right: 0; bottom: 0;
    background: rgba(11, 17, 32, 0.75);
    backdrop-filter: blur(6px);
    -webkit-backdrop-filter: blur(6px);
    z-index: 10001;
    display: none;
    justify-content: flex-end;
  }}
  .card-drawer-overlay.active {{
    display: flex;
  }}
  .card-drawer-panel {{
    width: 100%;
    max-width: 480px;
    height: 100%;
    background: var(--card);
    border-left: 1px solid var(--border);
    box-shadow: -10px 0 30px rgba(0,0,0,0.6);
    display: flex;
    flex-direction: column;
    animation: drawer-slide 0.25s cubic-bezier(0.16, 1, 0.3, 1);
  }}
  @keyframes drawer-slide {{
    from {{ transform: translateX(100%); }}
    to {{ transform: translateX(0); }}
  }}
  .card-drawer-header {{
    padding: 16px 20px;
    background: var(--raised);
    border-bottom: 1px solid var(--border);
    display: flex;
    align-items: center;
    justify-content: space-between;
  }}
  .card-drawer-body {{
    padding: 18px 20px;
    flex: 1;
    overflow-y: auto;
  }}

  {MT_CSS}
  {CITE_CSS}

  /* F-mi-dark: phone width */
  @media (max-width: 640px) {{
    body {{ padding: 16px 12px; }}
    .container {{ padding: 20px 14px; border-radius: 12px; }}
    details.rollup-box .rollup-content {{ padding: 14px 12px; }}
    .btn-toggle, .btn-pill {{ padding: 0 14px; }}
    .card-builder-trigger {{ bottom: 16px; right: 16px; padding: 9px 16px; font-size: 12.5px; }}
    .card-drawer-panel {{ max-width: 100%; }}
  }}
  @media (prefers-reduced-motion: reduce) {{
    *, *::before, *::after {{ animation: none !important; transition: none !important; scroll-behavior: auto !important; }}
  }}
  /* F-mi-dark: light print sheet so the PDF export stays printable */
  @media print {{
    :root {{
      color-scheme: light;
      --primary: #0f766e; --primary-light: #115e59; --secondary: #334155; --accent: #0f766e;
      --on-accent: #ffffff; --on-primary: #ffffff;
      --bg: #ffffff; --card: #ffffff; --raised: #f1f5f9; --border: #cbd5e1;
      --text: #0f172a; --body: #1e293b; --muted: #475569; --highlight: #f8fafc; --zebra: #f8fafc;
      --pos: #047857; --neg: #b91c1c; --warn: #b45309; --info: #0369a1; --star: #b45309;
    }}
    body {{ background: #ffffff; padding: 0; }}
    .container {{ border: 0; padding: 0; max-width: none; }}
    .badge, .pick-tier-badge {{ color: #0f172a !important; background: #f1f5f9 !important; border-color: #cbd5e1 !important; }}
    details.rollup-box[open] {{ box-shadow: none; }}
    details.rollup-box summary {{ min-height: 0; }}
    details.rollup-box > summary::after, .btn-toggle, .btn-pill, .rollup-controls, .global-rollup-bar, .back-to-top {{ display: none !important; }}
    .term-tooltip .tip-text {{ display: none !important; }}
    a {{ color: #0f766e; }}
  }}
</style>
</head>
<body data-week="{week_str}">
{"".join(outside_html)}
<div class="container">
{"".join(body_html)}
</div>
<script>
function toggleRollups(groupClass, expand) {{
  const selector = groupClass ? 'details.' + groupClass : 'details.rollup-box';
  document.querySelectorAll(selector).forEach(d => {{
    d.open = expand;
  }});
}}

function filterExperts(q) {{
  q = (q || '').trim().toLowerCase();
  let shownExperts = 0;
  document.querySelectorAll('details.show-grp').forEach(g => {{
    const sum = g.querySelector(':scope > summary');
    const showName = ((sum ? sum.textContent : '').split(' \u2014 ')[0] || '').toLowerCase();
    const showHit = !q || showName.includes(q);
    let any = false;
    g.querySelectorAll('details.section-4').forEach(e => {{
      const es = e.querySelector(':scope > summary');
      const hit = showHit || (es && es.textContent.toLowerCase().includes(q));
      e.style.display = hit ? '' : 'none';
      if (hit) {{ any = true; shownExperts++; }}
    }});
    const tbl = g.querySelector(':scope > .rollup-content > table');
    if (tbl) tbl.querySelectorAll('tbody tr').forEach(r => {{
      const c = r.cells && r.cells[0] ? r.cells[0].textContent.toLowerCase() : '';
      r.style.display = (showHit || c.includes(q)) ? '' : 'none';
    }});
    g.style.display = (any || showHit) ? '' : 'none';
    if (q) {{ if (g.dataset.prevOpen === undefined) g.dataset.prevOpen = g.open ? '1' : '0'; if (any || showHit) g.open = true; }}
    else if (g.dataset.prevOpen !== undefined) {{ g.open = g.dataset.prevOpen === '1'; delete g.dataset.prevOpen; }}
  }});
  const n = document.getElementById('expert-search-count');
  if (n) n.textContent = q ? (shownExperts + ' expert' + (shownExperts === 1 ? '' : 's') + ' found') : '';
}}

function toggleAllMainSections(expand) {{
  document.querySelectorAll('details.section-main').forEach(d => {{
    d.open = expand;
  }});
}}

function toggleAllRollups(expand) {{
  document.querySelectorAll('details.rollup-box').forEach(d => {{
    d.open = expand;
  }});
}}

function openTargetDetails(hash) {{
  if (!hash || hash === '#') return;
  if (window.__mtOnHash && window.__mtOnHash(hash)) return;
  const targetId = hash.replace('#', '');
  const targetEl = document.getElementById(targetId);
  if (targetEl && window.showTab) {{
    const pnl = targetEl.closest('.sc-panel');
    if (pnl) showTab(pnl.id.replace('sc-panel-', ''), false);
    if (targetId === 'sc-pick-sheet') {{ const pk0 = document.getElementById('pk-panel'); if (pk0) {{ pk0.classList.remove('collapsed'); document.body.classList.remove('pk-collapsed'); }} }}
  }}
  if (targetEl) {{
    let curr = targetEl;
    while (curr && curr !== document.body) {{
      if (curr.tagName && curr.tagName.toLowerCase() === 'details') {{
        curr.open = true;
      }}
      curr = curr.parentElement;
    }}
    if (targetEl.tagName && targetEl.tagName.toLowerCase() === 'details') {{
      targetEl.open = true;
    }}
    targetEl.classList.remove('target-pulse');
    void targetEl.offsetWidth;
    targetEl.classList.add('target-pulse');
    setTimeout(() => {{
      const topOffset = targetEl.getBoundingClientRect().top + window.pageYOffset - 90;
      window.scrollTo({{ top: Math.max(0, topOffset), behavior: 'auto' }});
    }}, 60);
  }}
}}

function parseCellValue(cell) {{
  if (!cell) return {{ type: 'empty', val: '' }};
  const ds = cell.querySelector('[data-sort]');
  if (ds) return {{ type: 'num', val: parseFloat(ds.getAttribute('data-sort')) }};
  const clone = cell.cloneNode(true);
  clone.querySelectorAll('.tip-text').forEach(el => el.remove());
  const txt = (clone.innerText || clone.textContent || '').trim();
  if (!txt) return {{ type: 'empty', val: '' }};

  // 1. Confidence Index (e.g. "96 / 100", "88 / 100")
  const mConf = txt.match(/(\\d+(?:\\.\\d+)?)\\s*\\/\\s*100/i);
  if (mConf) return {{ type: 'num', val: parseFloat(mConf[1]) }};

  // 2. Trench Mismatch SD (e.g. "+3.08 SD", "LAR +3.08 SD", "TB +2.10 SD")
  const mSD = txt.match(/([+-]?\\d+(?:\\.\\d+)?)\\s*SD/i);
  if (mSD) return {{ type: 'num', val: parseFloat(mSD[1]) }};

  // 3. Units (e.g. "4 Units", "2.0 Units", "0.5 Units", "4U")
  const mUnit = txt.match(/(\\d+(?:\\.\\d+)?)\\s*(?:Units?|U\\b)/i);
  if (mUnit) return {{ type: 'num', val: parseFloat(mUnit[1]) }};

  // 4. Source count consensus (e.g. "5-Source Consensus", "5-Show")
  const mShow = txt.match(/(\\d+)[-\\s](?:Source|Show)/i);
  if (mShow) return {{ type: 'num', val: parseFloat(mShow[1]) }};

  // 5. Market Edge with sign (e.g. "+5.4 on IND", "+2.2", "-4.6")
  const mEdge = txt.match(/([+-]\\d+(?:\\.\\d+)?)/);
  if (mEdge) return {{ type: 'num', val: parseFloat(mEdge[1]) }};

  // 6. Odds / Payout (e.g. "+765", "+645", "+371", "+150")
  const mOdds = txt.match(/[+](\\d{2,5})/);
  if (mOdds) return {{ type: 'num', val: parseFloat(mOdds[1]) }};

  // 7. General leading number (e.g. "49.1", "25.2", "-1.3")
  const mNum = txt.match(/^([+-]?\\d+(?:\\.\\d+)?)/);
  if (mNum) return {{ type: 'num', val: parseFloat(mNum[1]) }};

  return {{ type: 'str', val: txt.toLowerCase() }};
}}

function initSortableTables() {{
  document.querySelectorAll('table.sortable-table').forEach(table => {{
    const thead = table.querySelector('thead');
    if (!thead) return;
    const headerRow = thead.querySelector('tr');
    if (!headerRow) return;
    const headers = headerRow.querySelectorAll('th');
    const tbody = table.querySelector('tbody');
    if (!tbody) return;

    headers.forEach((th, colIdx) => {{
      th.classList.add('sortable-header');
      const sortIcon = document.createElement('span');
      sortIcon.className = 'sort-icon';
      sortIcon.setAttribute('aria-hidden', 'true');
      sortIcon.innerHTML = ' ↕';
      th.appendChild(sortIcon);

      th.addEventListener('click', (e) => {{
        if (e.target.closest('a')) return;

        const currentDir = th.getAttribute('data-sort-dir');
        let newDir = 'desc';
        if (!currentDir) {{
          const sampleRow = tbody.querySelector('tr');
          const sampleCell = sampleRow ? sampleRow.children[colIdx] : null;
          const parsed = parseCellValue(sampleCell);
          newDir = (parsed.type === 'num') ? 'desc' : 'asc';
        }} else {{
          newDir = currentDir === 'desc' ? 'asc' : 'desc';
        }}

        headers.forEach(h => {{
          h.removeAttribute('data-sort-dir');
          const icon = h.querySelector('.sort-icon');
          if (icon) icon.innerHTML = ' ↕';
        }});

        th.setAttribute('data-sort-dir', newDir);
        sortIcon.innerHTML = newDir === 'desc' ? ' ▼' : ' ▲';

        const rows = Array.from(tbody.querySelectorAll('tr'));

        rows.sort((rowA, rowB) => {{
          const cellA = rowA.children[colIdx];
          const cellB = rowB.children[colIdx];
          const valA = parseCellValue(cellA);
          const valB = parseCellValue(cellB);

          let cmp = 0;
          if (valA.type === 'num' && valB.type === 'num') {{
            cmp = valA.val - valB.val;
          }} else if (valA.type === 'num') {{
            cmp = 1;
          }} else if (valB.type === 'num') {{
            cmp = -1;
          }} else if (valA.type === 'empty' && valB.type === 'empty') {{
            cmp = 0;
          }} else if (valA.type === 'empty') {{
            cmp = -1;
          }} else if (valB.type === 'empty') {{
            cmp = 1;
          }} else {{
            cmp = String(valA.val).localeCompare(String(valB.val));
          }}

          return newDir === 'desc' ? -cmp : cmp;
        }});

        rows.forEach(r => tbody.appendChild(r));
      }});
    }});
  }});
}}

function initGroupSort() {{
  document.querySelectorAll('button.group-sort').forEach(btn => {{
    if (btn.dataset.ready) return;
    btn.dataset.ready = '1';
    btn.addEventListener('click', () => {{
      const sc = document.getElementById(btn.dataset.scope);
      if (!sc) return;
      const key = btn.dataset.key;
      const items = Array.from(sc.querySelectorAll(':scope > .sort-group'));
      items.sort((a, b) => key === 'conf' ? (parseFloat(b.dataset.conf) - parseFloat(a.dataset.conf)) : String(a.dataset[key]).localeCompare(String(b.dataset[key])));
      items.forEach(i => sc.appendChild(i));
      btn.parentElement.querySelectorAll('button.group-sort').forEach(b => b.classList.toggle('btn-primary', b === btn));
    }});
  }});
}}

function initSideTips() {{
  const tip = document.createElement('div');
  tip.className = 'side-tip';
  document.body.appendChild(tip);
  document.querySelectorAll('.side-toc a[data-tip]').forEach(a => {{
    a.addEventListener('mouseenter', () => {{
      const r = a.getBoundingClientRect();
      tip.textContent = a.dataset.tip;
      tip.style.display = 'block';
      const top = Math.min(r.top, window.innerHeight - tip.offsetHeight - 8);
      tip.style.top = Math.max(8, top) + 'px';
      tip.style.left = (r.right + 12) + 'px';
      if (r.right + 12 + tip.offsetWidth > window.innerWidth) {{ tip.style.left = Math.max(8, r.left) + 'px'; tip.style.top = (r.bottom + 6) + 'px'; }}
    }});
    a.addEventListener('mouseleave', () => {{ tip.style.display = 'none'; }});
  }});
}}

function initPickSheet() {{
  if (window.__pkInit) return;
  const boxes = Array.from(document.querySelectorAll('input.pk-box'));
  if (!boxes.length) return;
  window.__pkInit = true;
  const key = 'sc-dual:' + document.title;
  const WHO = ['andy', 'amanda'];
  let st = {{ andy: [], amanda: [], lock: [], frozen: null }};
  try {{ const j = JSON.parse(localStorage.getItem(key) || 'null'); if (j && typeof j === 'object') st = Object.assign(st, j); }} catch (e) {{}}
  const gameOf = {{}}; boxes.forEach(b => {{ gameOf[b.dataset.side] = b.dataset.game; }});
  const order = boxes.filter(b => b.dataset.who === 'andy').map(b => b.dataset.side);
  const byOrder = arr => arr.slice().sort((x, y) => order.indexOf(x) - order.indexOf(y));
  // one side per game: keep the first-listed side of any game that appears twice in a list
  const oneSide = arr => {{ const seen = {{}}; return arr.filter(sd => {{ const g = gameOf[sd]; if (g === undefined || seen[g]) return false; seen[g] = 1; return true; }}); }};
  WHO.forEach(w => {{ st[w] = oneSide(st[w] || []); }});
  const agreed = () => byOrder(st.andy.filter(s => st.amanda.includes(s)));
  const conflict = side => WHO.every(w => st[w].some(s => gameOf[s] === gameOf[side])) && !agreed().includes(side)
        && st.andy.some(s => gameOf[s] === gameOf[side] && !st.amanda.includes(s)) && st.amanda.some(s => gameOf[s] === gameOf[side] && !st.andy.includes(s));
  const save = () => {{ try {{ localStorage.setItem(key, JSON.stringify(st)); }} catch (e) {{}} }};
  const fill = (id, arr) => {{ const ol = document.getElementById(id); if (ol) ol.innerHTML = arr.length ? arr.map(s => '<li>' + s + '</li>').join('') : '<li class="pk-empty">none yet</li>'; }};
  const render = () => {{
    boxes.forEach(b => {{ b.checked = st[b.dataset.who].includes(b.dataset.side); b.disabled = !!st.frozen; }});
    const ag = agreed();
    WHO.forEach(w => {{
      fill('pk-list-' + w, byOrder(st[w]));
      const n = document.getElementById('pk-n-' + w); if (n) {{ n.textContent = st[w].length + '/5'; n.style.color = st[w].length === 5 ? '' : 'var(--warn)'; }}
    }});
    fill('pk-list-lock', byOrder(st.lock));
    const pc = document.getElementById('pk-counts'); if (pc) pc.textContent = 'A ' + st.andy.length + ' · M ' + st.amanda.length + ' · ✅ ' + ag.length + ' · 🔒 ' + st.lock.length;
    const nl = document.getElementById('pk-n-lock'); if (nl) {{ nl.textContent = st.lock.length + '/5'; nl.style.color = st.lock.length === 5 ? '' : 'var(--warn)'; }}
    document.querySelectorAll('.pk-st').forEach(sp => {{
      const sides = sp.dataset.game ? order.filter(x => gameOf[x] === sp.dataset.game) : [sp.dataset.side];
      const sh = x => (x || '').split(' ')[0];
      const lockS = sides.find(x => st.lock.includes(x)), agS = sides.find(x => st.andy.includes(x) && st.amanda.includes(x));
      const aS = sides.find(x => st.andy.includes(x)), mS = sides.find(x => st.amanda.includes(x));
      sp.textContent = lockS ? '🔒 ' + sh(lockS) : agS ? '✅ ' + sh(agS) : (aS && mS) ? '⚠️ conflict' : aS ? 'Andy ' + sh(aS) : mS ? 'Amanda ' + sh(mS) : '';
      const tr = sp.closest('tr'); if (tr) tr.style.background = lockS ? 'rgba(0,210,190,.18)' : agS ? 'rgba(0,210,190,.08)' : (aS && mS) ? 'rgba(251,191,36,.10)' : '';
    }});
    const only = WHO.map(w => st[w].filter(s => !ag.includes(s)).length);
    const stat = document.getElementById('pk-status');
    if (stat) stat.textContent = st.frozen ? '🔒 Frozen ' + st.frozen : (ag.length + ' agreed · Andy-only ' + only[0] + ' · Amanda-only ' + only[1]);
    const fz = document.getElementById('pk-freeze'); if (fz) fz.textContent = st.frozen ? 'Unfreeze' : 'Freeze official five';
    const fr = document.getElementById('pk-frozen'); if (fr) fr.textContent = st.frozen ? 'Locked in ' + st.frozen : (st.lock.length === 5 ? 'Ready to freeze' : '');
    save();
    (window.__pkHooks || []).forEach(f => {{ try {{ f(); }} catch (e) {{}} }});
  }};
  window.pkHas = (w, side) => st[w].includes(side);
  window.pkSidesOf = gid => order.filter(sd => gameOf[sd] === gid);
  window.pkToggle = (w, side) => {{
    if (st.frozen) return false;
    const arr = st[w];
    if (arr.includes(side)) st[w] = arr.filter(x => x !== side);
    else {{ st[w] = arr.filter(x => gameOf[x] !== gameOf[side]); st[w].push(side); }}
    st.lock = st.lock.filter(x => st.andy.includes(x) && st.amanda.includes(x));
    render(); return true;
  }};
  window.pkFrozen = () => !!st.frozen;
  boxes.forEach(b => b.addEventListener('change', () => {{
    if (st.frozen) return render();
    const arr = st[b.dataset.who], s = b.dataset.side;
    if (b.checked) {{ st[b.dataset.who] = arr.filter(x => gameOf[x] !== gameOf[s]); st[b.dataset.who].push(s); }}
    if (!b.checked) {{ st[b.dataset.who] = arr.filter(x => x !== s); }}
    st.lock = st.lock.filter(x => st.andy.includes(x) && st.amanda.includes(x));
    render();
  }}));
  window.pkSeed = w => {{ if (st.frozen) return; st[w] = boxes.filter(b => b.dataset.who === w && b.dataset.default === '1').map(b => b.dataset.side); render(); }};
  window.pkClear = w => {{ if (st.frozen) return; st[w] = []; st.lock = st.lock.filter(x => st.andy.includes(x) && st.amanda.includes(x)); render(); }};
  window.pkLockAgreed = () => {{ if (st.frozen) return; st.lock = agreed().slice(0, 5); render(); }};
  window.pkFreeze = () => {{
    if (st.frozen) {{ st.frozen = null; return render(); }}
    if (st.lock.length !== 5) {{ const s = document.getElementById('pk-status'); if (s) s.textContent = 'Need exactly 5 official picks (have ' + st.lock.length + ')'; return; }}
    st.frozen = new Date().toLocaleString(); render();
  }};
  window.copyPicks = () => {{
    const txt = byOrder(st.lock.length ? st.lock : agreed()).join('\\n');
    const done = () => {{ const c = document.getElementById('pk-status'); if (c) c.textContent = 'Copied official five'; }};
    if (navigator.clipboard) navigator.clipboard.writeText(txt).then(done).catch(() => {{ window.prompt && window.prompt('Copy:', txt); }});
    else if (window.prompt) window.prompt('Copy:', txt);
  }};
  const say = m => {{ const c = document.getElementById('pk-status'); if (c) c.textContent = m; }};
  const nm = w => w === 'andy' ? 'Andy' : 'Amanda';
  window.pkShare = w => {{
    const ta = document.getElementById('pk-text');
    let txt;
    if (w === 'sheet') txt = JSON.stringify(st);
    else txt = nm(w) + "'s SuperContest picks (" + st[w].length + "):\\n" + byOrder(st[w]).join('\\n');
    if (ta) {{ ta.value = txt; ta.focus(); ta.select(); }}
    const done = () => say(w === 'sheet' ? 'Full sheet copied' : nm(w) + "'s picks copied: paste them into a text or email");
    if (navigator.clipboard) navigator.clipboard.writeText(txt).then(done).catch(() => say('Select the text in the box and copy it'));
    else say('Select the text in the box and copy it');
  }};
  window.pkReceive = () => {{
    const ta = document.getElementById('pk-text'); const t = ta ? ta.value.trim() : '';
    if (!t) return say('Paste picks into the box first');
    if (st.frozen) return say('The sheet is frozen: unfreeze it before loading picks');
    if (t[0] === '{{') {{
      try {{ const j = JSON.parse(t); if (j && Array.isArray(j.andy) && Array.isArray(j.amanda)) {{ st = {{ andy: oneSide(j.andy), amanda: oneSide(j.amanda), lock: j.lock || [], frozen: j.frozen || null }}; render(); return say('Loaded the full sheet'); }} }} catch (e) {{}}
      return say('That sheet code did not load');
    }}
    const norm = t.replace(/[−–—]/g, '-');
    const who = /amanda/i.test(norm) && !/andy/i.test(norm) ? 'amanda' : /andy/i.test(norm) && !/amanda/i.test(norm) ? 'andy' : null;
    if (!who) return say('Start the text with Andy or Amanda so I know whose picks these are');
    const found = order.filter(side => {{ const re = new RegExp('(^|[^A-Za-z])' + side.replace(/[.+*?^${{}}()|[\\]\\\\]/g, '\\\\$&') + '(?![0-9.])', 'i'); return re.test(norm); }});
    if (!found.length) return say('No picks found: check the team and line (for example JAX -2.5)');
    const dupG = found.filter(sd => found.some(o => o !== sd && gameOf[o] === gameOf[sd]));
    if (dupG.length) return say('That text names both sides of a game (' + dupG.join(' and ') + '): you can only pick one side per game');
    st[who] = found; st.lock = st.lock.filter(x => st.andy.includes(x) && st.amanda.includes(x)); render();
    say('Loaded ' + nm(who) + "'s " + found.length + ' picks' + (found.length !== 5 ? ' (expected 5)' : ''));
  }};
  window.resetPicks = () => {{ st = {{ andy: [], amanda: [], lock: [], frozen: null }}; render(); }};
  render();
}}

function initTableFilters() {{
  document.querySelectorAll('.table-filter').forEach(bar => {{
    if (bar.dataset.ready) return;
    bar.dataset.ready = '1';
    const tables = () => {{
      if (bar.dataset.scope) {{
        const sc = document.getElementById(bar.dataset.scope);
        return sc ? Array.from(sc.querySelectorAll('table')) : [];
      }}
      let el = bar.nextElementSibling;
      while (el && el.tagName !== 'TABLE' && !el.querySelector('table')) el = el.nextElementSibling;
      return el ? [el.tagName === 'TABLE' ? el : el.querySelector('table')] : [];
    }};
    bar.querySelectorAll('button[data-filter]').forEach(btn => btn.addEventListener('click', () => {{
      bar.querySelectorAll('button[data-filter]').forEach(b => b.classList.toggle('btn-primary', b === btn));
      const want = btn.dataset.filter.toLowerCase();
      tables().forEach(t => {{
        const ths = Array.from(t.querySelectorAll('thead th'));
        const idx = ths.findIndex(th => {{
          const c = th.cloneNode(true);
          c.querySelectorAll('.tip-text,.sort-icon').forEach(e => e.remove());
          return c.textContent.trim().toLowerCase() === (bar.dataset.col || '').toLowerCase();
        }});
        let shown = 0;
        t.querySelectorAll('tbody tr').forEach(tr => {{
          const cell = idx >= 0 ? tr.children[idx] : null;
          let v = '';
          if (cell) {{ const cc = cell.cloneNode(true); cc.querySelectorAll('button, .bm-btn').forEach(e => e.remove()); v = cc.textContent.trim().toLowerCase(); }}
          const ok = want === 'all' || idx < 0 || v === want;
          tr.style.display = ok ? '' : 'none';
          if (ok) shown++;
        }});
        const grp = t.closest('.pool-group, .filter-group');
        if (grp) {{ grp.style.display = shown ? '' : 'none'; if (grp.tagName === 'DETAILS' && want !== 'all' && shown) grp.open = true; }}
      }});
      if (bar.dataset.scope) {{
        const sc2 = document.getElementById(bar.dataset.scope);
        if (sc2) sc2.querySelectorAll('.prop-cat-box').forEach(cb => {{
          const any = Array.from(cb.querySelectorAll('.filter-group')).some(g => g.style.display !== 'none');
          cb.style.display = any ? '' : 'none';
          const sm = cb.querySelector(':scope > summary');
          if (sm) {{
            if (!sm.dataset.orig) sm.dataset.orig = sm.textContent;
            const n = Array.from(cb.querySelectorAll('tbody tr')).filter(r => r.style.display !== 'none').length;
            sm.textContent = want === 'all' ? sm.dataset.orig : sm.dataset.orig.replace(/ — .*$/, '') + ' — ' + n + ' shown';
          }}
        }});
      }}
    }}));
  }});
}}

window.addEventListener('hashchange', () => openTargetDetails(location.hash));
const SC_TABS = [
  {{ id: 'home', label: '⭐ Home', anchors: ['sc-our-five'] }},
  {{ id: 'ranked', label: '📊 Ranked & Market', anchors: ['sc-ranked', 'sc-market'] }},
  {{ id: 'experts', label: '🧑‍🤝‍🧑 Experts', anchors: ['sc-experts'] }},
  {{ id: 'games', label: '🏈 Game-by-Game', anchors: ['sc-games'] }},
  {{ id: 'review', label: '📅 Season Review', anchors: ['sc-review'] }}
];
function initScTabs() {{
  if (window.__scTabs) return;
  const c = document.querySelector('.container');
  if (!c || !document.getElementById('sc-our-five') || !document.getElementById('pk-panel')) return;
  window.__scTabs = true;
  document.body.classList.add('sc-tabs-on');
  const kids = Array.from(c.children);
  const at = id => kids.findIndex(e => e.id === id);
  const start = at('sc-our-five');
  let end = at('disclaimer'); if (end < 0) end = kids.length;
  const pk = document.getElementById('pk-panel');
  const panels = {{}};
  SC_TABS.forEach(t => {{ const d = document.createElement('div'); d.className = 'sc-panel'; d.id = 'sc-panel-' + t.id; panels[t.id] = d; }});
  // which tab does each anchor start?
  const marks = [];
  SC_TABS.forEach(t => {{ t.anchors.forEach(a => {{ const i = at(a); if (i >= 0) marks.push([i, t.id]); }}); }});
  marks.sort((x, y) => x[0] - y[0]);
  const tabFor = i => {{ let cur = 'home'; marks.forEach(m => {{ if (i >= m[0]) cur = m[1]; }}); return cur; }};
  const nav = document.createElement('header'); nav.className = 'sc-topbar'; nav.id = 'sc-tabs';
  const h1 = c.querySelector('h1'); const brand = document.createElement('span'); brand.className = 'sc-brand'; brand.textContent = h1 ? h1.textContent.replace(/NFL |Intelligence Report/g, '').replace(/\s+/g, ' ').trim() : 'SuperContest';
  const whoBox = document.createElement('div'); whoBox.className = 'sc-who'; whoBox.id = 'sc-who';
  whoBox.innerHTML = '<span class="sc-who-l">I am</span><button type="button" data-w="andy">Andy</button><button type="button" data-w="amanda">Amanda</button>';
  nav.appendChild(whoBox);
  const setWho = (w) => {{
    document.body.dataset.who = w || '';
    whoBox.querySelectorAll('button').forEach(b => b.classList.toggle('on', b.dataset.w === w));
    try {{ if (w) localStorage.setItem('sc-who', w); else localStorage.removeItem('sc-who'); }} catch (e) {{}}
    (window.__pkHooks || []).forEach(f => {{ try {{ f(); }} catch (e) {{}} }});
  }};
  window.scSetWho = setWho;
  whoBox.querySelectorAll('button').forEach(b => b.addEventListener('click', () => setWho(document.body.dataset.who === b.dataset.w ? '' : b.dataset.w)));
  try {{ const w0 = localStorage.getItem('sc-who'); if (w0 === 'andy' || w0 === 'amanda') {{ document.body.dataset.who = w0; whoBox.querySelectorAll('button').forEach(b => b.classList.toggle('on', b.dataset.w === w0)); }} }} catch (e) {{}}
  nav.appendChild(brand);
  const menu = document.createElement('div'); menu.id = 'sc-menu'; menu.className = 'sc-menu'; menu.hidden = true; document.body.appendChild(menu);
  const clean = t => (t || '').replace(/\s+/g, ' ').trim();
  const sumText = el => {{ const sm = el && el.querySelector(':scope > summary'); return sm ? clean(sm.innerText || sm.textContent) : ''; }};
  const itemsFor = id => {{
    const P = panels[id]; const out = [];
    const add = (label, target, sub) => {{ if (label && document.getElementById(target)) out.push({{ label: label, target: target, sub: !!sub }}); }};
    if (id === 'home') {{
      add('⭐ Our five', 'sc-grp-five');
      P.querySelectorAll('#sc-grp-five .sc-five').forEach(d => add(sumText(d).split(' — ')[0], d.id, true));
      add('🔁 Alternates', 'sc-grp-alt');
      P.querySelectorAll('#sc-grp-alt .sc-five').forEach(d => add(sumText(d).split(' — ')[0], d.id, true));
      add('🚨 Contest status', 'sc-status');
    }} else if (id === 'ranked') {{
      add('📊 1. Every side, ranked', 'sc-ranked'); add('📖 How the ranking works', 'sc-score-defs', true); add('⚖️ 2. Contest lines vs. the market', 'sc-market'); add('📊 Betting splits by game', 'sc-splits', true);
    }} else if (id === 'experts') {{
      P.querySelectorAll('details.sc-exp').forEach(d => add(sumText(d).split(':')[0], d.id));
    }} else if (id === 'games') {{
      P.querySelectorAll('details[id]').forEach(d => {{
        const m = d.id.match(/^sc-game-([a-z]+)-([a-z]+)$/); if (!m) return;
        const t = sumText(d); const ko = (t.match(/(Mon|Tue|Wed|Thu|Fri|Sat|Sun)\s+\d{{1,2}}:\d{{2}}/) || [''])[0];
        add((t.indexOf('⭐') > -1 ? '⭐ ' : '') + m[1].toUpperCase() + ' @ ' + m[2].toUpperCase() + (ko ? ' · ' + ko : ''), d.id);
      }});
    }} else if (id === 'review') {{
      add('📋 How the SuperContest works', 'sc-rules'); add('📅 5. Season review & rules of thumb', 'sc-review');
      P.querySelectorAll('details[id]').forEach(d => {{ if (d.id === 'sc-rules' || /^section-/.test(d.id) || /-experts$/.test(d.id)) return; add(sumText(d), d.id, true); }});
    }}
    return out;
  }};
  const closeMenu = () => {{ menu.hidden = true; nav.querySelectorAll('.sc-caret').forEach(b => b.setAttribute('aria-expanded', 'false')); }};
  const openMenu = (id, btn) => {{
    const items = itemsFor(id);
    if (!items.length) return closeMenu();
    menu.innerHTML = items.map(i => '<a class="sc-mi' + (i.sub ? ' sub' : '') + '" href="#' + i.target + '">' + i.label.replace(/</g, '&lt;') + '</a>').join('');
    const r = btn.getBoundingClientRect();
    menu.hidden = false;
    const pr = pk.classList.contains('collapsed') ? null : pk.getBoundingClientRect();
    const limit = (pr && pr.top < 5 && pr.width < window.innerWidth * 0.6) ? pr.left : window.innerWidth;
    const w = Math.min(340, limit - 16);
    menu.style.width = w + 'px';
    menu.style.left = Math.max(8, Math.min(r.left, limit - w - 8)) + 'px';
    menu.style.top = (r.bottom + 4) + 'px';
    menu.style.maxHeight = (window.innerHeight - r.bottom - 16) + 'px';
    nav.querySelectorAll('.sc-caret').forEach(b => b.setAttribute('aria-expanded', b === btn ? 'true' : 'false'));
  }};
  menu.addEventListener('click', ev => {{ if (ev.target.closest('.sc-mi')) closeMenu(); }});
  document.addEventListener('click', ev => {{ if (!ev.target.closest('.sc-menu') && !ev.target.closest('.sc-tabwrap')) closeMenu(); }});
  document.addEventListener('keydown', ev => {{ if (ev.key === 'Escape') closeMenu(); }});
  window.addEventListener('scroll', closeMenu, {{ passive: true }});
  SC_TABS.forEach(t => {{
    const w = document.createElement('div'); w.className = 'sc-tabwrap'; w.dataset.tab = t.id;
    const b = document.createElement('button'); b.type = 'button'; b.className = 'sc-tab'; b.dataset.tab = t.id; b.textContent = t.label;
    b.addEventListener('click', () => {{ closeMenu(); showTab(t.id, true); }});
    const cb = document.createElement('button'); cb.type = 'button'; cb.className = 'sc-caret'; cb.setAttribute('aria-label', 'Sections in ' + t.label); cb.setAttribute('aria-expanded', 'false'); cb.textContent = '▾';
    cb.addEventListener('click', ev => {{ ev.stopPropagation(); if (!menu.hidden && cb.getAttribute('aria-expanded') === 'true') return closeMenu(); showTab(t.id, false); openMenu(t.id, w); cb.setAttribute('aria-expanded', 'true'); }});
    w.appendChild(b); w.appendChild(cb); nav.appendChild(w);
  }});
  const tools = document.createElement('div'); tools.className = 'sc-tools';
  [['⊞', 'Expand all sections', true], ['⊟', 'Collapse all sections', false]].forEach(x => {{ const tb = document.createElement('button'); tb.type = 'button'; tb.className = 'sc-tool'; tb.textContent = x[0]; tb.title = x[1]; tb.setAttribute('aria-label', x[1]); tb.addEventListener('click', () => {{ closeMenu(); if (typeof toggleAllRollups === 'function') toggleAllRollups(x[2]); }}); tools.appendChild(tb); }});
  nav.appendChild(tools);
  const wrap = document.createElement('div'); wrap.className = 'sc-panels';
  const first = kids[start];
  c.insertBefore(wrap, first);
  const sideToc = document.getElementById('side-toc'); if (sideToc) sideToc.remove();
  const tocBtn = document.querySelector('.toc-toggle'); if (tocBtn) tocBtn.remove();
  document.body.insertBefore(nav, document.body.firstChild);
  SC_TABS.forEach(t => wrap.appendChild(panels[t.id]));
  for (let i = 0; i < kids.length; i++) {{
    const e = kids[i];
    if (e === pk || e.id === 'sc-pick-sheet') continue;
    if (e.id === 'sc-rules') {{ panels.review.insertBefore(e, panels.review.firstChild); continue; }}
    if (e.id === 'sc-status') {{ panels.home.appendChild(e); continue; }}
    if (i >= start && i < end) panels[tabFor(i)].appendChild(e);
  }}
  const stBox = document.getElementById('sc-status'); if (stBox) panels.home.appendChild(stBox);
  const rb = document.getElementById('section-rec-box');
  if (rb && rb.parentNode) {{
    const rc = rb.querySelector(':scope > .rollup-content');
    const prev = rb.previousElementSibling;
    if (prev && prev.classList.contains('rollup-controls')) prev.remove();
    if (rc) {{ while (rc.firstChild) rb.parentNode.insertBefore(rc.firstChild, rb); }}
    rb.remove();
  }}
  document.body.appendChild(pk);
  window.showTab = (id, scroll) => {{
    if (!panels[id]) id = 'home';
    SC_TABS.forEach(t => {{ panels[t.id].classList.toggle('active', t.id === id); }});
    nav.querySelectorAll('.sc-tab').forEach(b => b.classList.toggle('active', b.dataset.tab === id));
    nav.querySelectorAll('.sc-tabwrap').forEach(w => w.classList.toggle('active', w.dataset.tab === id));
    try {{ sessionStorage.setItem('sc-tab', id); }} catch (e) {{}}
    if (scroll) window.scrollTo({{ top: Math.max(0, wrap.getBoundingClientRect().top + window.pageYOffset - nav.offsetHeight - 10), behavior: 'smooth' }});
  }};
  const route = () => {{
    const h = (location.hash || '').replace('#', '');
    if (h === 'sc-pick-sheet') {{ document.body.classList.remove('pk-collapsed'); pk.classList.remove('collapsed'); return; }}
    const el = h ? document.getElementById(h) : null;
    const pnl = el ? el.closest('.sc-panel') : null;
    if (pnl) {{ showTab(pnl.id.replace('sc-panel-', ''), false); setTimeout(() => openTargetDetails(location.hash), 0); }}
  }};
  let saved = 'home'; try {{ saved = sessionStorage.getItem('sc-tab') || 'home'; }} catch (e) {{}}
  showTab(saved, false);
  window.addEventListener('hashchange', route);
  if (location.hash) route();
  window.pkTogglePanel = () => {{ const cl = pk.classList.toggle('collapsed'); document.body.classList.toggle('pk-collapsed', cl); }};
  if (window.innerWidth < 1200) {{ pk.classList.add('collapsed'); document.body.classList.add('pk-collapsed'); }}
}}

{MT_JS}
window.addEventListener('DOMContentLoaded', () => {{
  initScTabs();
  initMasterTabs();
  if (location.hash) {{
    openTargetDetails(location.hash);
  }}
  initSortableTables();
  initTableFilters();
  initGroupSort();
  initSideTips();
  initPickSheet();
  initBookmarks();
  initScCardButtons();
}});
if (document.readyState !== 'loading') {{
  initScTabs();
  initMasterTabs();
  initSortableTables();
  initTableFilters();
  initGroupSort();
  initSideTips();
  initPickSheet();
  initBookmarks();
  initScCardButtons();
}}

function initScCardButtons() {{
  if (window.__scCardInit) return;
  const boxes = Array.from(document.querySelectorAll('input.pk-box'));
  if (!boxes.length || !window.pkToggle) return;
  window.__scCardInit = true;
  const gids = Array.from(new Set(boxes.map(b => b.dataset.game)));
  const flat = t => (t || '').replace(/[\\u2212\\u2013\\u2014]/g, '-').replace(/\\s+/g, ' ');
  const esc = t => t.replace(/[.+*?^${{}}()|[\\]\\\\]/g, '\\\\$&');
  const who = () => document.body.dataset.who || '';
  const say = m => {{ const c = document.getElementById('pk-status'); if (c) c.textContent = m; }};
  const info = [];
  document.querySelectorAll('.container .bm-btn').forEach(btn => {{
    const row = btn.closest('tr, .bm-item, .at-glance-item, details');
    if (!row) return;
    const slots = row.querySelectorAll('.sc-slot');
    if (slots.length) {{
      btn.remove();
      slots.forEach(sl => {{
        if (sl.dataset.done) return; sl.dataset.done = '1';
        const sd = sl.dataset.side; const wrap = document.createElement('span'); wrap.className = 'sc-cb one';
        const lab = document.createElement('label'); const cb = document.createElement('input'); cb.type = 'checkbox'; cb.dataset.side = sd; lab.appendChild(cb); wrap.appendChild(lab); sl.appendChild(wrap);
        cb.title = 'Tick to add ' + sd + ' to the pick sheet for whoever is chosen at the top left';
        cb.addEventListener('click', ev => ev.stopPropagation());
        cb.addEventListener('change', () => {{
          if (!who()) {{
            cb.checked = false;
            const wb = document.getElementById('sc-who'); if (wb) {{ wb.classList.add('flash'); setTimeout(() => wb.classList.remove('flash'), 1600); }}
            return say('Choose Andy or Amanda at the top left first');
          }}
          if (window.pkFrozen()) {{ cb.checked = window.pkHas(who(), sd); return say('The sheet is frozen: unfreeze it to change picks'); }}
          if (cb.checked !== window.pkHas(who(), sd)) window.pkToggle(who(), sd);
        }});
        info.push({{ boxes: [cb] }});
      }});
      return;
    }}
    if (btn.closest('details.sc-exp')) {{ btn.remove(); return; }}
    const t = flat(row.tagName === 'TR' ? Array.from(row.children).map(c => c.textContent).join(' | ') : row.textContent);
    const hit = gids.filter(g => {{ const m = g.split('@'); return new RegExp('(^|[^A-Z])' + m[0] + ' ?@ ?' + m[1] + '([^A-Z]|$)').test(t); }});
    let gid = hit.length === 1 ? hit[0] : null;
    let sides = [];
    if (!gid) {{
      const named = boxes.filter(b => b.dataset.who === 'andy' && new RegExp('(^|[^A-Za-z])' + esc(b.dataset.side) + '(?![0-9.])').test(t)).map(b => b.dataset.side);
      const gs = Array.from(new Set(named.map(sd => boxes.find(b => b.dataset.side === sd).dataset.game)));
      if (gs.length === 1) {{ gid = gs[0]; if (named.length === 1) sides = named; }}
    }}
    if (!gid) {{ btn.remove(); return; }}
    const all = window.pkSidesOf(gid);
    if (!sides.length) {{
      const named = all.filter(sd => new RegExp('(^|[^A-Za-z])' + esc(sd) + '(?![0-9.])').test(t));
      const th = row.closest('table') && row.closest('table').querySelector('thead');
      const sideRow = !!th && Array.from(th.querySelectorAll('th')).some(h => h.textContent.trim().indexOf('Side') === 0);
      sides = (sideRow && named.length === 1) ? named : all;
    }}
    const wrap = document.createElement('span'); wrap.className = 'sc-cb' + (sides.length === 1 ? ' one' : '');
    wrap.title = 'Tick to add to the pick sheet for whoever is chosen at the top left';
    const boxesHere = sides.map(sd => {{
      const lab = document.createElement('label'); const cb = document.createElement('input'); cb.type = 'checkbox'; cb.dataset.side = sd;
      lab.appendChild(cb); if (sides.length > 1) lab.appendChild(document.createTextNode(sd)); wrap.appendChild(lab);
      cb.addEventListener('click', ev => ev.stopPropagation());
      cb.addEventListener('change', () => {{
        if (!who()) {{
          cb.checked = false;
          const wb = document.getElementById('sc-who'); if (wb) {{ wb.classList.add('flash'); setTimeout(() => wb.classList.remove('flash'), 1600); }}
          return say('Choose Andy or Amanda at the top left first');
        }}
        if (window.pkFrozen()) {{ cb.checked = window.pkHas(who(), sd); return say('The sheet is frozen: unfreeze it to change picks'); }}
        if (cb.checked !== window.pkHas(who(), sd)) window.pkToggle(who(), sd);
      }});
      return cb;
    }});
    btn.replaceWith(wrap);
    info.push({{ boxes: boxesHere }});
  }});
  const refresh = () => info.forEach(i => i.boxes.forEach(cb => {{
    const w = who(); cb.checked = !!w && window.pkHas(w, cb.dataset.side);
    cb.parentElement.classList.toggle('on', cb.checked);
  }}));
  (window.__pkHooks = window.__pkHooks || []).push(refresh);
  refresh();
}}

let modalHistory = [];

function pushModalState() {{
  const overlay = document.getElementById('matchup-modal-overlay');
  const titleEl = document.getElementById('modal-matchup-title');
  const bodyEl = document.getElementById('modal-matchup-body');
  const jumpLink = document.getElementById('modal-jump-link');
  if (overlay && overlay.classList.contains('active') && bodyEl) {{
    modalHistory.push({{
      title: titleEl.innerHTML,
      body: bodyEl.innerHTML,
      jumpHref: jumpLink.getAttribute('href'),
      scrollTop: bodyEl.scrollTop
    }});
  }}
}}

function popModalState() {{
  if (modalHistory.length === 0) return false;
  const prev = modalHistory.pop();
  const titleEl = document.getElementById('modal-matchup-title');
  const bodyEl = document.getElementById('modal-matchup-body');
  const jumpLink = document.getElementById('modal-jump-link');
  const backBtn = document.getElementById('modal-back-btn');
  
  titleEl.innerHTML = prev.title;
  bodyEl.innerHTML = prev.body;
  jumpLink.setAttribute('href', prev.jumpHref);
  bodyEl.scrollTop = prev.scrollTop;
  
  if (backBtn) {{
    backBtn.style.display = modalHistory.length > 0 ? 'inline-flex' : 'none';
  }}
  return true;
}}

function openMatchupModal(gid, isPush = false) {{
  const gm = String(gid).match(/game-[a-z]+-[a-z]+/i);
  const cleanId = gm ? gm[0].toLowerCase() : String(gid).replace(/^#[\\/]?/, '').replace(/^.*[\\/]([^\\/]+)$/, '$1').replace(/\\.html$/, '');
  const boxId = cleanId.endsWith('-box') ? cleanId : (cleanId + '-box');
  const box = document.getElementById(boxId) ||
              document.getElementById(cleanId) ||
              document.querySelector('details#' + boxId) ||
              document.querySelector('details#' + cleanId) ||
              document.querySelector('[data-page="' + cleanId + '"]');
  if (!box) return false;
  const overlay = document.getElementById('matchup-modal-overlay');
  const titleEl = document.getElementById('modal-matchup-title');
  const bodyEl = document.getElementById('modal-matchup-body');
  const jumpLink = document.getElementById('modal-jump-link');
  const backBtn = document.getElementById('modal-back-btn');
  if (!overlay || !bodyEl) return false;
  
  if (isPush) {{
    pushModalState();
  }} else {{
    modalHistory = [];
  }}
  
  const summ = box.querySelector('summary') || box.querySelector('h2') || box.querySelector('.page-h');
  const content = box.querySelector('.rollup-content') || box;
  
  titleEl.innerHTML = summ ? (summ.querySelector('.sum-text')?.innerHTML || summ.innerHTML) : 'Matchup Intel';
  bodyEl.innerHTML = content.innerHTML;
  jumpLink.setAttribute('href', '#' + (box.id || cleanId));
  bodyEl.scrollTop = 0;
  
  bodyEl.querySelectorAll('details').forEach(d => d.open = true);
  
  if (backBtn) {{
    backBtn.style.display = modalHistory.length > 0 ? 'inline-flex' : 'none';
  }}
  
  overlay.classList.add('active');
  document.body.style.overflow = 'hidden';
  return true;
}}

function closeMatchupModal() {{
  const overlay = document.getElementById('matchup-modal-overlay');
  if (overlay) {{
    overlay.classList.remove('active');
    document.body.style.overflow = '';
  }}
  modalHistory = [];
  const backBtn = document.getElementById('modal-back-btn');
  if (backBtn) backBtn.style.display = 'none';
}}

function initBookmarks() {{
  const wk = document.body.getAttribute('data-week') || '{week_str}';
  const KEY = 'pr-bookmarks-w' + wk;
  function loadBms() {{
    try {{
      const a = JSON.parse(localStorage.getItem(KEY) || '[]');
      return Array.isArray(a) ? a : [];
    }} catch (e) {{ return []; }}
  }}
  function saveBms(a) {{
    try {{ localStorage.setItem(KEY, JSON.stringify(a)); }} catch (e) {{}}
  }}
  function hashStr(s) {{
    let x = 5381;
    for (let i = 0; i < s.length; i++) {{ x = ((x << 5) + x + s.charCodeAt(i)) | 0; }}
    return (x >>> 0).toString(36);
  }}
  function cleanTxt(t) {{ return (t || '').replace(/\\s+/g, ' ').trim(); }}
  function updateCount() {{
    const n = loadBms().length;
    document.querySelectorAll('.bm-count').forEach(c => {{ c.textContent = n ? String(n) : ''; }});
  }}
  function renderDrawer() {{
    const list = document.getElementById('drawer-bm-list') || document.getElementById('bm-list');
    if (!list) return;
    const items = loadBms();
    if (!items.length) {{
      list.innerHTML = '<p class=\"muted-note\">No picks bookmarked yet. Tap the ☆ star next to any pick, prop, or table row to build your weekend card.</p>';
      return;
    }}
    const groups = {{}};
    const orderG = [];
    items.forEach(b => {{
      const cat = b.ctx || 'Starred Picks';
      if (!groups[cat]) {{ groups[cat] = []; orderG.push(cat); }}
      groups[cat].push(b);
    }});
    list.innerHTML = '';
    orderG.forEach(cat => {{
      const g = document.createElement('div');
      g.className = 'bm-group';
      const h = document.createElement('h3');
      h.textContent = cat;
      g.appendChild(h);
      groups[cat].forEach(b => {{
        const r = document.createElement('div');
        r.className = 'bm-row';
        const l = document.createElement('a');
        l.href = '#' + b.id;
        l.innerHTML = '<strong>' + cleanTxt(b.text.slice(0, 160)) + '</strong>';
        l.addEventListener('click', (ev) => {{
          ev.preventDefault();
          closeDrawer();
          const target = document.getElementById(b.id);
          if (target) {{
            if (window.__mtOnHash && window.__mtOnHash('#' + b.id)) {{ target.classList.add('bm-flash'); return; }}
            let q = target;
            while (q) {{ if (q.tagName === 'DETAILS') q.open = true; q = q.parentElement; }}
            target.scrollIntoView({{ block: 'center', behavior: 'smooth' }});
            target.classList.add('bm-flash');
          }}
        }});
        const x = document.createElement('button');
        x.type = 'button';
        x.className = 'btn-toggle';
        x.style.padding = '4px 10px';
        x.style.fontSize = '12px';
        x.textContent = '✕ Remove';
        x.addEventListener('click', () => {{
          saveBms(loadBms().filter(z => z.id !== b.id));
          const btn = document.querySelector('#' + b.id + ' .bm-btn');
          if (btn) {{ btn.classList.remove('on'); btn.setAttribute('aria-pressed', 'false'); }}
          updateCount();
          renderDrawer();
        }});
        r.appendChild(l);
        r.appendChild(x);
        g.appendChild(r);
      }});
      list.appendChild(g);
    }});
  }}
  
  // Wire table rows and items
  const marks = {{}};
  loadBms().forEach(b => {{ marks[b.id] = 1; }});
  const rows = document.querySelectorAll('.container table tbody tr, .container .bm-item, .container details.leg, .container details.section-tickets');
  rows.forEach(el => {{
    if (el.closest('.no-bm') || el.closest('#bm-list') || el.querySelector('.bm-btn') || el.dataset.bmReady) return;
    // only real picks get a star: table rows must sit in a table whose header names a Pick / Prop / Leg column
    // (not trends, news, market boards or the expert registry, whose rows are context, not bets)
    if (el.matches('tr')) {{
      const tb = el.closest('table');
      const heads = tb ? Array.from(tb.querySelectorAll('thead th')).map(th => {{ const c = th.cloneNode(true); c.querySelectorAll('.tip-text,.sort-icon,.bm-btn').forEach(e => e.remove()); return (c.textContent || '').trim().toLowerCase(); }}) : [];
      const okTable = heads.some(h => /^(pick|prop|leg)(\b|$)/.test(h)) && !heads.includes('trend') && !heads.includes('rationale');
      if (!okTable || el.closest('#mt-panel-experts, #mt-panel-market, #mt-panel-about, .game-narrative, [id$="-trends"], [id$="-news"]')) return;
    }}
    el.dataset.bmReady = '1';
    const text = cleanTxt(el.textContent);
    if (!text || text.length < 3) return;
    const host = el.matches('tr') ? el.querySelector('td') : (el.querySelector('.bm-host') || el.querySelector('summary') || el);
    if (!host) return;
    if (!el.id) {{ let hid = 'bm-' + hashStr(text.slice(0, 160)), k = 1; while (document.getElementById(hid)) hid = 'bm-' + hashStr(text.slice(0, 160)) + '-' + (k++); el.id = hid; }}
    
    const b = document.createElement('button');
    b.type = 'button';
    b.className = 'bm-btn' + (marks[el.id] ? ' on' : '');
    b.setAttribute('aria-label', 'Add pick to card');
    b.setAttribute('aria-pressed', marks[el.id] ? 'true' : 'false');
    b.title = 'Add pick to card';
    b.innerHTML = '<span class="bm-icon">' + (marks[el.id] ? '★' : '☆') + '</span><span class="bm-text">' + (marks[el.id] ? 'On Card' : '+ Card') + '</span>';
    b.addEventListener('click', (ev) => {{
      ev.preventDefault();
      ev.stopPropagation();
      const a = loadBms();
      const idx = a.findIndex(z => z.id === el.id);
      if (idx >= 0) {{
        a.splice(idx, 1);
        b.classList.remove('on');
        b.setAttribute('aria-pressed', 'false');
        b.innerHTML = '<span class="bm-icon">☆</span><span class="bm-text">+ Card</span>';
      }} else {{
        let ctx = '';
        const parentBox = el.closest('details');
        const summ = parentBox?.querySelector('summary');
        if (summ) ctx = cleanTxt(summ.textContent).slice(0, 70);
        const pageSec = el.closest('[data-page]');
        const pg = pageSec ? pageSec.getAttribute('data-page') : 'index';
        const title = pageSec ? (pageSec.getAttribute('data-title') || pg) : '';
        a.push({{ id: el.id, page: pg, title: title, text: text.slice(0, 180), ctx: ctx, t: Date.now() }});
        b.classList.add('on');
        b.setAttribute('aria-pressed', 'true');
        b.innerHTML = '<span class="bm-icon">★</span><span class="bm-text">On Card</span>';
      }}
      saveBms(a);
      updateCount();
      renderDrawer();
    }});
    host.insertBefore(b, host.firstChild);
  }});
  updateCount();
  renderDrawer();

  // Wire drawer controls
  const drawerBtn = document.getElementById('card-builder-btn');
  const drawerOverlay = document.getElementById('card-drawer-overlay');
  const drawerClose = document.getElementById('card-drawer-close');
  const bmClear = document.getElementById('drawer-bm-clear') || document.getElementById('bm-clear');
  const bmCopy = document.getElementById('drawer-bm-copy') || document.getElementById('bm-copy');

  function openDrawer() {{
    if (drawerOverlay) {{
      renderDrawer();
      drawerOverlay.classList.add('active');
      document.body.style.overflow = 'hidden';
    }}
  }}
  function closeDrawer() {{
    if (drawerOverlay) {{
      drawerOverlay.classList.remove('active');
      document.body.style.overflow = '';
    }}
  }}

  if (drawerBtn) drawerBtn.addEventListener('click', openDrawer);
  if (drawerClose) drawerClose.addEventListener('click', closeDrawer);
  if (drawerOverlay) drawerOverlay.addEventListener('click', (e) => {{
    if (e.target === drawerOverlay) closeDrawer();
  }});
  if (bmClear) bmClear.addEventListener('click', () => {{
    saveBms([]);
    document.querySelectorAll('.bm-btn.on').forEach(b => {{ b.classList.remove('on'); b.setAttribute('aria-pressed', 'false'); }});
    updateCount();
    renderDrawer();
  }});
  if (bmCopy) bmCopy.addEventListener('click', () => {{
    const a = loadBms();
    const txt = a.map(b => '- ' + b.text + (b.ctx ? ' (' + b.ctx + ')' : '')).join('\\n');
    if (navigator.clipboard) {{
      navigator.clipboard.writeText(txt).then(() => {{
        bmCopy.textContent = 'Copied!';
        setTimeout(() => {{ bmCopy.textContent = 'Copy Card to Clipboard'; }}, 1600);
      }}).catch(() => {{ window.prompt && window.prompt('Copy Card:', txt); }});
    }} else if (window.prompt) {{
      window.prompt('Copy Card:', txt);
    }}
  }});
}}

document.addEventListener('click', (e) => {{
  // 1. External links always open in a new tab/window
  const ext = e.target.closest('a[href^="http://"], a[href^="https://"]');
  if (ext) {{
    ext.setAttribute('target', '_blank');
    ext.setAttribute('rel', 'noopener noreferrer');
    return;
  }}

  // 2. Modal header controls
  if (e.target.closest('#modal-close-btn')) {{
    closeMatchupModal();
    return;
  }}
  if (e.target.closest('#modal-back-btn')) {{
    popModalState();
    return;
  }}
  const overlay = document.getElementById('matchup-modal-overlay');
  if (e.target === overlay) {{
    closeMatchupModal();
    return;
  }}
  const jumpLink = e.target.closest('#modal-jump-link');
  if (jumpLink) {{
    e.preventDefault();
    const hash = jumpLink.getAttribute('href');
    closeMatchupModal();
    openTargetDetails(hash);
    return;
  }}

  // 2b. Ranked-table matchup links open the Master Intel pop-up (fall back to the Game-by-Game jump)
  const miLink = e.target.closest('a[data-mi]');
  if (miLink && typeof openMatchupModal === 'function' && openMatchupModal(miLink.getAttribute('data-mi'), false) !== false) {{
    e.preventDefault(); e.stopPropagation(); return;
  }}

  // 3. Game write-up / matchup links (ANYWHERE on page or inside modal)
  const link = e.target.closest('a[href]');
  if (link) {{
    // Do NOT intercept clicks on header nav, page nav, or matchup game cards
    if (link.closest('.site-header') || link.classList.contains('game-card') || link.closest('.page-nav')) return;

    const href = link.getAttribute('href') || '';
    const text = (link.textContent || '').trim().toLowerCase();
    
    // Check if this is a game write-up / matchup link
    const isPill = link.classList.contains('at-glance-pill');
    const isWriteup = text.includes('write-up') || text.includes('full matchup') || isPill || link.classList.contains('table-link');
    const gm = href.match(/(game-[a-z]+-[a-z]+)/i);
    if (gm && isWriteup && !href.startsWith('#sc-')) {{
      e.preventDefault();
      const insideModal = link.closest('#modal-matchup-body');
      if (openMatchupModal(gm[1], !!insideModal) === false && href.indexOf('.html') > -1) {{ window.location.href = href; }}
      return;
    }}

    // Intra-modal navigation for other elements ("pop-up on a pop-up")
    const insideModal = link.closest('#modal-matchup-body');
    if (insideModal) {{
      if (href.startsWith('#props-')) {{
        const cleanP = href.replace(/^#/, '');
        const propBox = document.getElementById(cleanP) || document.querySelector('[id*="' + cleanP + '"]');
        if (propBox) {{
          e.preventDefault();
          pushModalState();
          const titleEl = document.getElementById('modal-matchup-title');
          const bodyEl = document.getElementById('modal-matchup-body');
          const jumpLink = document.getElementById('modal-jump-link');
          const backBtn = document.getElementById('modal-back-btn');
          const summ = propBox.querySelector('summary') || propBox.closest('details')?.querySelector('summary');
          const content = propBox.querySelector('.rollup-content') || propBox.closest('details')?.querySelector('.rollup-content') || propBox;
          titleEl.innerHTML = summ ? summ.innerHTML : 'Player Props';
          bodyEl.innerHTML = content.innerHTML;
          jumpLink.setAttribute('href', href);
          bodyEl.scrollTop = 0;
          bodyEl.querySelectorAll('details').forEach(d => d.open = true);
          if (backBtn) backBtn.style.display = 'inline-flex';
          return;
        }}
      }}
      if (href.startsWith('#exp-')) {{
        const cleanExp = href.replace(/^#/, '');
        const expBox = document.getElementById(cleanExp) || document.getElementById('box-' + cleanExp) || document.querySelector('[id*="' + cleanExp + '"]');
        if (expBox) {{
          e.preventDefault();
          pushModalState();
          const titleEl = document.getElementById('modal-matchup-title');
          const bodyEl = document.getElementById('modal-matchup-body');
          const jumpLink = document.getElementById('modal-jump-link');
          const backBtn = document.getElementById('modal-back-btn');
          const summ = expBox.querySelector('summary') || expBox.closest('details')?.querySelector('summary');
          const content = expBox.querySelector('.rollup-content') || expBox.closest('details')?.querySelector('.rollup-content') || expBox;
          titleEl.innerHTML = summ ? summ.innerHTML : 'Expert Pick Card';
          bodyEl.innerHTML = content.innerHTML;
          jumpLink.setAttribute('href', href);
          bodyEl.scrollTop = 0;
          bodyEl.querySelectorAll('details').forEach(d => d.open = true);
          if (backBtn) backBtn.style.display = 'inline-flex';
          return;
        }}
      }}
      // If clicking other internal anchor inside modal, don't scroll background
      e.preventDefault();
      return;
    }}
    
    // Other page anchor navigation (ignore hash routes like #/props or #/bookmarks)
    if (href.startsWith('#') && !href.startsWith('#/') && href !== '#' && href !== '#executive-board') {{
      e.preventDefault();
      history.pushState(null, '', href);
      openTargetDetails(href);
    }}
  }}
}});

document.addEventListener('keydown', (e) => {{
  if (e.key === 'Escape') {{
    closeMatchupModal();
    const drawer = document.getElementById('card-drawer-overlay');
    if (drawer && drawer.classList.contains('active')) {{
      drawer.classList.remove('active');
      document.body.style.overflow = '';
    }}
  }}
}});
</script>
<div id="matchup-modal-overlay" class="matchup-modal-overlay" role="dialog" aria-modal="true">
  <div class="matchup-modal-dialog">
    <div class="matchup-modal-header">
      <div style="display:flex;align-items:center;gap:10px;flex:1;min-width:0;">
        <button id="modal-back-btn" class="btn-toggle" style="display:none;padding:5px 12px;font-size:12px;font-weight:700;" type="button">← Back</button>
        <div id="modal-matchup-title" class="matchup-modal-title"></div>
      </div>
      <div class="matchup-modal-actions">
        <a id="modal-jump-link" class="btn-toggle" href="#">Scroll to Section ↓</a>
        <button id="modal-close-btn" class="btn-toggle btn-primary" type="button">✕ Close (Esc)</button>
      </div>
    </div>
    <div id="modal-matchup-body" class="matchup-modal-body"></div>
  </div>
</div>
<button id="card-builder-btn" class="card-builder-trigger" type="button" aria-label="Open Card Builder">
  📋 My Card <span class="bm-count"></span>
</button>

<div id="card-drawer-overlay" class="card-drawer-overlay" role="dialog" aria-modal="true">
  <div class="card-drawer-panel">
    <div class="card-drawer-header">
      <div style="display:flex;align-items:center;gap:10px;">
        <span style="font-size:20px;">📋</span>
        <h3 style="margin:0;font-size:16px;font-weight:800;color:var(--text);">Weekend Card Builder</h3>
        <span class="bm-count" style="display:inline-block;"></span>
      </div>
      <button id="card-drawer-close" class="btn-toggle btn-primary" type="button">✕ Close (Esc)</button>
    </div>
    <div class="card-drawer-body">
      <p class="muted-note" style="margin-top:0;">Picks, legs, props, and table rows you starred across this report. Saved in this browser for reviewing and placing your weekend card.</p>
      <div class="rollup-controls" style="margin-bottom:14px;">
        <button class="btn-toggle btn-primary" id="drawer-bm-copy" type="button">Copy Card to Clipboard</button>
        <button class="btn-toggle" id="drawer-bm-clear" type="button">Clear Card</button>
      </div>
      <div id="drawer-bm-list">
        <p class="muted-note">No picks bookmarked yet. Tap the ☆ star next to any pick, prop, or table row to build your weekend card.</p>
      </div>
    </div>
  </div>
</div>
</body>
</html>"""

    # Post-process full_html:
    # 1. Enforce target="_blank" rel="noopener noreferrer" on all external links
    full_html = re.sub(r'<a\s+href="(https?://[^"]+)"(?![^>]*target=)', r'<a href="\1" target="_blank" rel="noopener noreferrer"', full_html)

    # 2. Bold / heavy visual styling for Confidence Stars
    def _rep_star(m):
        full = m.group(0)
        num_str = m.group(1)
        stars_str = m.group(2)
        try:
            val = float(num_str)
        except ValueError:
            return full
        if val >= 4.0:
            tier_class = "star-tier-high"
        elif val >= 3.0:
            tier_class = "star-tier-mid"
        else:
            tier_class = "star-tier-low"
        
        active_stars = ''.join([c for c in stars_str if c == '★'])
        empty_stars = ''.join([c for c in stars_str if c == '☆'])
        
        return (f'<span class="star-rating {tier_class}" title="{val:.1f} stars">'
                f'<span class="star-num">{val:.1f}</span> '
                f'<span class="stars-active">{active_stars}</span>'
                f'<span class="stars-muted">{empty_stars}</span>'
                f'</span>')

    full_html = re.sub(r'(\d\.\d)\s+([★☆]{5})', _rep_star, full_html)

    Path(html_path).write_text(full_html, encoding='utf-8')
    print(f"Successfully generated HTML: {html_path}")

if __name__ == '__main__':
    if len(sys.argv) >= 4:
        md = sys.argv[1]
        docx_path = sys.argv[2]
        html_path = sys.argv[3]
    else:
        md = 'E:/dev/projects/NFL_Dashboard/scratch/nfl_week1_supercontest_intelligence_summary.md'
        docx_path = 'E:/dev/projects/NFL_Dashboard/scratch/nfl_week1_supercontest_intelligence_summary.docx'
        html_path = 'E:/dev/projects/NFL_Dashboard/scratch/nfl_week1_supercontest_intelligence_summary.html'
    generate_docx(md, docx_path)
    generate_html(md, html_path)
